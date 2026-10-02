from __future__ import annotations

from dataclasses import FrozenInstanceError
import csv
from hashlib import sha256
import json
import math
from pathlib import Path

import openpyxl
import pytest

import dm2c_m23_canonical_release as release_module
from dm2c_m23_canonical import (
    APPLICATION_CLASSES,
    OPTIONAL_CONTEXT_PREDICATES,
    PRINCIPAL_PREDICATES,
    CanonicalLPGGraph,
    stable_id,
)
from dm2c_m23_canonical_release import (
    CaseConfig,
    CaseConfigError,
    EvidenceAssembly,
    ReleaseArtifacts,
    ReleaseBuildError,
    assemble_material_evidence,
    build_validation_rows,
    build_release,
    load_case_config,
)
from dm2c_m23_calculation import ValidationIssue, ValidationResultSet, validate_candidates
from dm2c_m23_ifc import CanonicalIFCExtractor
from dm2c_multigranular_carbon_kg import FactorLibrary
from dm2c_canonical_v2_reader import load_canonical_v2_context
from dm2c_case_study_facts import build_construction_facts


ROOT = Path(__file__).resolve().parents[1]
ACTUAL_CONFIG = ROOT / "inputs" / "case_study" / "m23_canonical_v2_case.json"
TYPEA1_CONFIG = ROOT / "inputs" / "case_study" / "m23_canonical_v2_case_typea1.json"
ACTUAL_IFC = ROOT / "typed_completed.ifc"
TYPEA1_IFC = ROOT / "typea1_full.ifc"
ONTOLOGY = ROOT / "mic-carbon-ontology.ttl"
EXPECTED_IFC_HASH = "9934A332500BCBDA7EA40835803546982926F99B91CC84F229D59721807B937B"
EXPECTED_WORKBOOK_HASH = "60A71D2FCAAD3A38CD7D8B86A9175E346D035DDC7B29EE3DEE2E7E34F69C3E06"
ARTIFACT_NAMES = {
    "multigranular_carbon_kg.json",
    "multigranular_carbon_kg.cypher",
    "multigranular_carbon_kg_stats.json",
    "multigranular_carbon_kg_validation.csv",
    "m2_alignment_report.json",
    "case_version_manifest.json",
}
OUTPUT_NAMES = {
    "graphJson": "multigranular_carbon_kg.json",
    "cypher": "multigranular_carbon_kg.cypher",
    "stats": "multigranular_carbon_kg_stats.json",
    "validation": "multigranular_carbon_kg_validation.csv",
    "alignmentReport": "m2_alignment_report.json",
}


def _hash_size(path: Path) -> tuple[str, int]:
    payload = path.read_bytes()
    return sha256(payload).hexdigest().upper(), len(payload)


def _binding(path: Path) -> dict[str, object]:
    digest, size = _hash_size(path)
    return {"path": str(path.resolve()), "sha256": digest, "sizeBytes": size}


@pytest.fixture(scope="session")
def actual_extraction():
    return CanonicalIFCExtractor(ACTUAL_IFC).extract()


def _write_factor_workbook(path: Path) -> str:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Embodied Carbon Coefficients"
    sheet.append(
        [
            "Material Category",
            "Material Sub-Type",
            "Selected A1-A3 Factor",
            "Unit",
            "Original Source / Database",
            "Factor Year",
            "Life-Cycle Boundary",
            "Source URL",
            "Geography",
            "Proxy Status",
        ]
    )
    sheet.append(
        [
            "Controlled material",
            "Fixture row",
            0.5,
            "kgCO2e/kg",
            "controlled fixture source",
            "2026",
            "A1-A3",
            "doi:controlled-factor",
            "Hong Kong",
            "primary",
        ]
    )
    workbook.save(path)
    digest, _ = _hash_size(path)
    return f"sha256:{digest}:Embodied Carbon Coefficients:2"


def _write_energy_factor_workbook(path: Path) -> dict[str, str]:
    workbook = openpyxl.Workbook()
    material_sheet = workbook.active
    material_sheet.title = "Embodied Carbon Coefficients"
    material_sheet.append(
        [
            "Material Category",
            "Material Sub-Type",
            "Selected A1-A3 Factor",
            "Unit",
            "Original Source / Database",
            "Factor Year",
            "Life-Cycle Boundary",
            "Source URL",
            "Geography",
            "Proxy Status",
        ]
    )
    material_sheet.append(
        [
            "Controlled material",
            "Fixture row",
            0.5,
            "kgCO2e/kg",
            "controlled fixture source",
            "2026",
            "A1-A3",
            "doi:controlled-factor",
            "Hong Kong",
            "primary",
        ]
    )
    energy_sheet = workbook.create_sheet("Energy Emission Factors")
    energy_sheet.append(
        [
            "Category",
            "Energy Type",
            "Selected Factor",
            "Unit",
            "Factor Year",
            "Original Source / Database",
            "Life-Cycle Boundary",
            "Source URL",
            "Geography",
            "Proxy Status",
        ]
    )
    energy_sheet.append(
        [
            "Electricity",
            "Guangdong factory proxy",
            0.62,
            "kgCO2e/kWh",
            "2026",
            "controlled energy source",
            "A1-A3",
            "doi:controlled-electricity",
            "Guangdong",
            "proxy",
        ]
    )
    energy_sheet.append(
        [
            "Fuel",
            "Diesel lifecycle",
            3.2,
            "kgCO2e/L",
            "2026",
            "controlled diesel source",
            "A1-A3",
            "doi:controlled-diesel",
            "Hong Kong",
            "proxy",
        ]
    )
    workbook.save(path)
    digest, _ = _hash_size(path)
    return {
        "electricity": f"sha256:{digest}:Energy Emission Factors:2",
        "diesel": f"sha256:{digest}:Energy Emission Factors:3",
    }


@pytest.fixture
def controlled_case(tmp_path: Path, actual_extraction):
    inputs = tmp_path / "controlled-inputs"
    inputs.mkdir()
    factor_path = inputs / "controlled-factors.xlsx"
    factor_source_row_id = _write_factor_workbook(factor_path)

    quantities_by_component: dict[str, list] = {}
    for quantity in actual_extraction.design_quantities:
        quantities_by_component.setdefault(quantity.component_id, []).append(quantity)
    association = next(
        row
        for row in actual_extraction.material_associations
        if any(
            quantity.normalized_unit == "m3"
            for quantity in quantities_by_component.get(row.component_id, ())
        )
    )
    quantity = next(
        row
        for row in quantities_by_component[association.component_id]
        if row.normalized_unit == "m3"
    )
    rejected_association = next(
        row
        for row in actual_extraction.material_associations
        if row.id != association.id
    )
    other_component = next(
        row
        for row in actual_extraction.components
        if row.id != rejected_association.component_id
    )

    valid_record = {
        "recordId": "controlled-material:accepted",
        "associationId": association.id,
        "componentId": association.component_id,
        "materialId": association.material_id,
        "factorSourceRowId": factor_source_row_id,
        "formulaCode": "volume_density",
        "operands": [
            {
                "operandId": "operand:controlled-volume",
                "role": "material_volume",
                "value": quantity.source_value,
                "unit": quantity.source_unit,
                "sourceId": quantity.id,
                "designQuantityId": quantity.id,
            },
            {
                "operandId": "operand:controlled-density",
                "role": "density",
                "value": 2700.0,
                "unit": "kg/m3",
                "sourceId": "reviewed-density:controlled:1",
                "designQuantityId": None,
            },
        ],
        "measured": False,
    }
    rejected_record = {
        **valid_record,
        "recordId": "controlled-material:rejected-ownership",
        "associationId": rejected_association.id,
        "componentId": other_component.id,
        "materialId": rejected_association.material_id,
    }
    evidence_path = inputs / "controlled-material-evidence.json"
    evidence_path.write_text(
        json.dumps(
            {
                "schemaVersion": "m23-canonical-v2",
                "records": [valid_record, rejected_record],
            },
            ensure_ascii=False,
            sort_keys=True,
            indent=2,
        ),
        encoding="utf-8",
    )

    config_payload = {
        "schemaVersion": "m23-canonical-v2",
        "releaseProfile": "controlled-fixture",
        "releaseReady": True,
        "ifc": _binding(ACTUAL_IFC),
        "factorWorkbook": _binding(factor_path),
        "ontology": _binding(ONTOLOGY),
        "materialEvidence": _binding(evidence_path),
        "module": {
            "sourceIdentity": "controlled-module:1",
            "name": "Controlled module",
            "identitySource": "controlled_fixture",
        },
        "requestedScope": "A1-A3",
        "includeOpenings": False,
        "factoryInput": None,
        "generatedAtUtc": "2026-07-20T00:00:00Z",
    }
    config_path = inputs / "controlled-case.json"
    config_path.write_text(
        json.dumps(config_payload, ensure_ascii=False, sort_keys=True, indent=2),
        encoding="utf-8",
    )
    return {
        "configPath": config_path,
        "configPayload": config_payload,
        "factorPath": factor_path,
        "factorSourceRowId": factor_source_row_id,
        "evidencePath": evidence_path,
        "association": association,
        "rejectedAssociation": rejected_association,
        "quantity": quantity,
        "validRecord": valid_record,
        "rejectedRecord": rejected_record,
    }


def _read_manifest(release_dir: Path) -> dict:
    return json.loads((release_dir / "case_version_manifest.json").read_text("utf-8"))


def _assert_no_staging_dirs(out_root: Path, release_id: str) -> None:
    if out_root.exists():
        assert not list(out_root.glob(f".{release_id}.staging-*"))


def test_actual_case_config_builds_backbone_only_release_without_material_evidence(
    tmp_path: Path,
):
    payload = json.loads(ACTUAL_CONFIG.read_text(encoding="utf-8"))
    factor_path = (ACTUAL_CONFIG.parent / payload["factorWorkbook"]["path"]).resolve()
    payload["ifc"] = _binding(ACTUAL_IFC)
    payload["factorWorkbook"] = _binding(factor_path)
    payload["ontology"] = _binding(ONTOLOGY)
    payload["factoryInput"] = None
    payload["factoryTargetMap"] = None
    config_path = tmp_path / "actual-backbone-only-case.json"
    config_path.write_text(
        json.dumps(payload, ensure_ascii=False, sort_keys=True, indent=2),
        encoding="utf-8",
    )
    config = load_case_config(config_path)

    assert isinstance(config, CaseConfig)
    assert config.schema_version == "m23-canonical-v2"
    assert config.release_profile == "actual-case"
    assert config.release_ready is False
    assert config.ifc.sha256 == EXPECTED_IFC_HASH
    assert config.ifc.size_bytes == 30_277_955
    assert config.factor_workbook.sha256 == EXPECTED_WORKBOOK_HASH
    assert config.factor_workbook.size_bytes == 30_242
    assert config.material_evidence is None
    assert config.factory_input is None
    assert config.factory_target_map is None
    assert config.requested_scope == "A1-A3"

    out_root = tmp_path / "actual-release"
    release_dir = build_release(config, out_root, "actual_case_backbone_only")
    manifest = _read_manifest(release_dir)
    stats = json.loads(
        (release_dir / "multigranular_carbon_kg_stats.json").read_text("utf-8")
    )
    graph = json.loads(
        (release_dir / "multigranular_carbon_kg.json").read_text("utf-8")
    )

    assert {path.name for path in release_dir.iterdir()} == ARTIFACT_NAMES
    assert manifest["releaseProfile"] == "actual-case"
    assert manifest["releaseReady"] is False
    assert manifest["inputs"]["materialEvidence"] is None
    assert manifest["coverage"] == {
        "acceptedByKind": {},
        "acceptedCount": 0,
        "candidateCount": 0,
        "coverageFormula": "acceptedCount / candidateCount",
        "coverageValue": 0.0,
        "rejectedByKind": {},
        "rejectedByReason": {},
        "rejectedCount": 0,
    }
    assert stats["calculationCounts"] == {"energy": 0, "material": 0, "validZero": 0}
    assert stats["applicationClassCounts"]["BuildingComponent"] > 0
    assert stats["applicationClassCounts"]["CarbonEmission"] == 0
    assert all(
        "CarbonEmission" not in node["labels"]
        and "MaterialConsumption" not in node["labels"]
        and "EnergyConsumption" not in node["labels"]
        for node in graph["nodes"]
    )
    assert (release_dir / "m2_alignment_report.json").exists()
    _assert_no_staging_dirs(out_root, "actual_case_backbone_only")

    context = load_canonical_v2_context(
        release_dir, allow_unready_backbone=True
    )
    structural = build_construction_facts(context)
    assert set(structural) == {"release", "module", "graph", "ifc_composition"}
    assert structural["graph"]["class_counts"]["CarbonEmission"] == 0
    assert structural["graph"]["class_counts"]["MaterialConsumption"] == 0
    assert structural["ifc_composition"]["component_count"] == structural["graph"][
        "class_counts"
    ]["BuildingComponent"]


def test_typea1_case_config_is_a_backbone_only_actual_case():
    """The second-model configuration must never trigger model-specific accounting."""
    config = load_case_config(TYPEA1_CONFIG)

    assert config.release_profile == "actual-case"
    assert config.release_ready is False
    assert config.ifc.resolved_path == TYPEA1_IFC
    assert config.ifc.sha256 == "0DC4C1525217A12650BEF064048F38A1891F5C34827D1B18AE6D053260C92C9D"
    assert config.ifc.size_bytes == 32_387_972
    assert config.material_evidence is None
    assert config.factory_input is None
    assert config.factory_target_map is None
    assert config.factor_workbook.resolved_path == (
        ACTUAL_CONFIG.parent
        / json.loads(ACTUAL_CONFIG.read_text(encoding="utf-8"))["factorWorkbook"]["path"]
    ).resolve()
    assert config.ontology.resolved_path == ONTOLOGY


def test_actual_case_factory_input_populates_process_context_and_energy_facts(
    tmp_path: Path,
    actual_extraction,
):
    inputs = tmp_path / "factory-inputs"
    inputs.mkdir()
    factor_path = inputs / "energy-factors.xlsx"
    factor_rows = _write_energy_factor_workbook(factor_path)
    component = next(row for row in actual_extraction.components if row.global_id)
    factory_path = inputs / "factory-data.json"
    factory_path.write_text(
        json.dumps(
            {
                "p0_energy_records": [
                    {
                        "record_id": "SYN_E_DIRECT_COMPONENT_001",
                        "source_id": "source:synthetic_component_meter",
                        "target_id": component.global_id,
                        "target_level": "component",
                        "stage": "Component assembly",
                        "activity": "Component-level fixture energy",
                        "batch_id": "BATCH_DIRECT",
                        "energy_carrier": "electricity",
                        "quantity_value": "10",
                        "quantity_unit": "kWh",
                        "metering_level": "workstation",
                        "production_line": "LINE_A",
                        "workstation_id": "WS_01",
                        "period_start": "2026-01-01",
                        "period_end": "2026-01-31",
                        "record_role": "direct_component_log",
                        "defect_flags": "",
                    },
                    {
                        "record_id": "SYN_E_SHARED_DIESEL_001",
                        "source_id": "source:synthetic_factory_meter",
                        "target_id": "unknown",
                        "target_level": "unknown",
                        "stage": "Factory support",
                        "activity": "Shared factory diesel",
                        "batch_id": "ALL_BATCHES",
                        "energy_carrier": "diesel",
                        "quantity_value": "2",
                        "quantity_unit": "L",
                        "metering_level": "factory_month",
                        "production_line": "",
                        "workstation_id": "",
                        "period_start": "2026-01-01",
                        "period_end": "2026-01-31",
                        "record_role": "shared_factory_meter",
                        "defect_flags": "",
                    },
                    {
                        "record_id": "SYN_E_PROCESS_ONLY_001",
                        "source_id": "source:synthetic_line_meter",
                        "target_id": "unknown",
                        "target_level": "unknown",
                        "stage": "Factory support",
                        "activity": "Unresolved shared line electricity",
                        "batch_id": "",
                        "energy_carrier": "electricity",
                        "quantity_value": "5",
                        "quantity_unit": "kWh",
                        "metering_level": "line",
                        "production_line": "LINE_Z",
                        "workstation_id": "",
                        "period_start": "2026-01-01",
                        "period_end": "2026-01-31",
                        "record_role": "unresolved_shared_line_meter",
                        "defect_flags": "",
                    },
                ],
                "p0_allocation_rows": [
                    {
                        "record_id": "SYN_E_SHARED_DIESEL_001",
                        "target_id": "Main_Modularization_Model",
                        "target_level": "module",
                        "batch_id": "BATCH_1",
                        "allocation_basis": "mass",
                        "allocation_value": "100",
                        "allocation_unit": "kg",
                        "allocation_fraction": "0.25",
                        "module_id": "BATCH_1",
                    },
                    {
                        "record_id": "SYN_E_SHARED_DIESEL_001",
                        "target_id": "FactoryModule_002",
                        "target_level": "module",
                        "batch_id": "BATCH_2",
                        "allocation_basis": "mass",
                        "allocation_value": "100",
                        "allocation_unit": "kg",
                        "allocation_fraction": "0.25",
                        "module_id": "BATCH_2",
                    },
                    {
                        "record_id": "SYN_E_SHARED_DIESEL_001",
                        "target_id": "FactoryModule_003",
                        "target_level": "module",
                        "batch_id": "BATCH_3",
                        "allocation_basis": "mass",
                        "allocation_value": "100",
                        "allocation_unit": "kg",
                        "allocation_fraction": "0.25",
                        "module_id": "BATCH_3",
                    },
                    {
                        "record_id": "SYN_E_SHARED_DIESEL_001",
                        "target_id": "FactoryModule_004",
                        "target_level": "module",
                        "batch_id": "BATCH_4",
                        "allocation_basis": "mass",
                        "allocation_value": "100",
                        "allocation_unit": "kg",
                        "allocation_fraction": "0.25",
                        "module_id": "BATCH_4",
                    },
                ],
                "p0_source_metadata": [
                    {
                        "source_id": "source:synthetic_component_meter",
                        "source_name": "Synthetic component meter",
                        "source_type": "synthetic_factory_energy_record",
                    },
                    {
                        "source_id": "source:synthetic_factory_meter",
                        "source_name": "Synthetic shared factory meter",
                        "source_type": "synthetic_factory_energy_record",
                    },
                    {
                        "source_id": "source:synthetic_line_meter",
                        "source_name": "Synthetic line meter",
                        "source_type": "synthetic_factory_energy_record",
                    },
                ],
            },
            ensure_ascii=False,
            sort_keys=True,
            indent=2,
        ),
        encoding="utf-8",
    )
    config_payload = {
        "schemaVersion": "m23-canonical-v2",
        "releaseProfile": "actual-case",
        "releaseReady": False,
        "ifc": _binding(ACTUAL_IFC),
        "factorWorkbook": _binding(factor_path),
        "ontology": _binding(ONTOLOGY),
        "materialEvidence": None,
        "module": {
            "sourceIdentity": "Main_Modularization_Model",
            "name": "Main Modularization Model",
            "identitySource": "case_manifest",
        },
        "requestedScope": "A1-A3",
        "includeOpenings": False,
        "factoryInput": _binding(factory_path),
        "generatedAtUtc": "2026-07-21T00:00:00Z",
    }
    config_path = inputs / "actual-factory-case.json"
    config_path.write_text(
        json.dumps(config_payload, ensure_ascii=False, sort_keys=True, indent=2),
        encoding="utf-8",
    )

    release_dir = build_release(
        load_case_config(config_path),
        tmp_path / "factory-release",
        "actual_factory_context",
    )
    graph = json.loads((release_dir / "multigranular_carbon_kg.json").read_text("utf-8"))
    stats = json.loads(
        (release_dir / "multigranular_carbon_kg_stats.json").read_text("utf-8")
    )
    manifest = _read_manifest(release_dir)

    assert manifest["inputs"]["factoryInput"]["sha256"] == _binding(factory_path)["sha256"]
    assert manifest["syntheticFactoryInputsUsed"] is True
    assert stats["applicationClassCounts"]["EnergyConsumption"] == 3
    assert stats["applicationClassCounts"]["CarbonEmission"] == 3
    assert stats["applicationClassCounts"]["ManufacturingProcessTemplate"] == 3
    assert stats["applicationClassCounts"]["ProductionStage"] == 13
    assert stats["applicationClassCounts"]["ManufacturingActivity"] == 69
    assert stats["applicationClassCounts"]["ManufacturingResource"] > 0
    assert stats["applicationClassCounts"]["ProductionBatch"] == 1
    assert stats["relationCounts"]["hasProcessTemplate"] > 0
    assert stats["relationCounts"]["hasStage"] == 13
    assert stats["relationCounts"]["hasActivity"] == 69
    assert stats["relationCounts"]["usesResource"] >= 69
    assert stats["relationCounts"]["directlyPrecedes"] > 0
    assert stats["validation"]["acceptedByKind"] == {"energy": 3}
    assert stats["validation"]["acceptedCount"] == 3
    assert stats["attributionCounts"] == {
        "allocated": 1,
        "direct": 1,
        "process_only": 1,
    }
    assert stats["calculationCounts"] == {"energy": 3, "material": 0, "validZero": 0}
    process_labels = {
        "ManufacturingProcessTemplate",
        "ProductionStage",
        "ManufacturingActivity",
        "ManufacturingResource",
    }
    serialized_process = json.dumps(
        [
            node.get("props", {})
            for node in graph["nodes"]
            if process_labels.intersection(node.get("labels", ()))
        ],
        ensure_ascii=False,
    )
    assert "镀锌" not in serialized_process
    assert "存放发运" not in serialized_process
    assert "叉车" not in serialized_process
    direct_node = next(
        node
        for node in graph["nodes"]
        if node["props"].get("recordId") == "SYN_E_DIRECT_COMPONENT_001"
        and "EnergyConsumption" in node["labels"]
    )
    allocated_node = next(
        node
        for node in graph["nodes"]
        if node["props"].get("recordId") == "SYN_E_SHARED_DIESEL_001"
        and "EnergyConsumption" in node["labels"]
    )
    process_node = next(
        node
        for node in graph["nodes"]
        if node["props"].get("recordId") == "SYN_E_PROCESS_ONLY_001"
        and "EnergyConsumption" in node["labels"]
    )
    assert direct_node["props"]["measured"] is False
    assert allocated_node["props"]["measured"] is False
    assert process_node["props"]["dataProvenance"] == "synthetic"
    assert direct_node["props"]["factorSourceId"] == factor_rows["electricity"]
    assert allocated_node["props"]["factorSourceId"] == factor_rows["diesel"]
    direct_edges = [
        edge
        for edge in graph["edges"]
        if edge["src"] == direct_node["id"] and edge["type"] == "recordedForObject"
    ]
    allocated_edges = [
        edge
        for edge in graph["edges"]
        if edge["src"] == allocated_node["id"] and edge["type"] == "recordedForObject"
    ]
    process_edges = [
        edge
        for edge in graph["edges"]
        if edge["src"] == process_node["id"] and edge["type"] == "recordedForObject"
    ]
    assert len(direct_edges) == 1
    assert direct_edges[0]["tgt"] == component.id
    direct_process_edges = [
        edge
        for edge in graph["edges"]
        if edge["src"] == direct_node["id"] and edge["type"] == "associatedWithProcess"
    ]
    direct_resource_edges = [
        edge
        for edge in graph["edges"]
        if edge["src"] == direct_node["id"] and edge["type"] == "recordedForResource"
    ]
    assert len(direct_process_edges) == 1
    assert len(direct_resource_edges) == 1
    direct_activity_id = direct_process_edges[0]["tgt"]
    direct_resource_id = direct_resource_edges[0]["tgt"]
    direct_activity = next(node for node in graph["nodes"] if node["id"] == direct_activity_id)
    direct_resource = next(node for node in graph["nodes"] if node["id"] == direct_resource_id)
    assert direct_activity["labels"] == ["ManufacturingActivity"]
    assert direct_resource["labels"] == ["ManufacturingResource"]
    stage_edge = next(
        edge
        for edge in graph["edges"]
        if edge["type"] == "hasActivity" and edge["tgt"] == direct_activity_id
    )
    resource_edge = next(
        edge
        for edge in graph["edges"]
        if edge["type"] == "usesResource"
        and edge["src"] == direct_activity_id
        and edge["tgt"] == direct_resource_id
    )
    stage = next(node for node in graph["nodes"] if node["id"] == stage_edge["src"])
    assert stage["labels"] == ["ProductionStage"]
    template_edge = next(
        edge
        for edge in graph["edges"]
        if edge["type"] == "hasStage" and edge["tgt"] == stage["id"]
    )
    template = next(node for node in graph["nodes"] if node["id"] == template_edge["src"])
    assert template["labels"] == ["ManufacturingProcessTemplate"]
    assert resource_edge["props"]["activityId"] == direct_activity["props"]["activityId"]
    assert len(allocated_edges) > 1
    assert all(edge["props"]["allocated"] is True for edge in allocated_edges)
    assert {
        "allocationBasis",
        "allocatedFraction",
    } <= set(allocated_edges[0]["props"])
    assert allocated_node["props"]["unattributedFraction"] == pytest.approx(0.75)
    assert math.isclose(
        math.fsum(edge["props"]["allocatedFraction"] for edge in allocated_edges),
        0.25,
        rel_tol=0.0,
        abs_tol=1e-9,
    )
    assert math.isclose(
        math.fsum(edge["props"]["allocatedFraction"] for edge in allocated_edges)
        + allocated_node["props"]["unattributedFraction"],
        1.0,
        rel_tol=0.0,
        abs_tol=1e-9,
    )
    assert all(
        "BuildingComponent"
        in next(node["labels"] for node in graph["nodes"] if node["id"] == edge["tgt"])
        for edge in allocated_edges
    )
    assert process_edges == []


def test_release_remaps_legacy_factory_target_without_graph_mapping_metadata(
    tmp_path: Path,
    actual_extraction,
):
    inputs = tmp_path / "factory-target-map-inputs"
    inputs.mkdir()
    factor_path = inputs / "energy-factors.xlsx"
    _write_energy_factor_workbook(factor_path)
    component = actual_extraction.components[0]
    legacy_global_id = "legacy-component-global-id"
    factory_path = inputs / "factory-data.json"
    factory_path.write_text(
        json.dumps(
            {
                "p0_energy_records": [
                    {
                        "record_id": "MAP_REMAP_DIRECT_COMPONENT_001",
                        "source_id": "source:legacy_component_meter",
                        "target_id": legacy_global_id,
                        "target_level": "component",
                        "stage": "Component assembly",
                        "activity": "Mapped component energy",
                        "batch_id": "",
                        "energy_carrier": "electricity",
                        "quantity_value": "10",
                        "quantity_unit": "kWh",
                        "metering_level": "workstation",
                        "production_line": "LINE_MAP",
                        "workstation_id": "WS_MAP",
                        "period_start": "2026-01-01",
                        "period_end": "2026-01-31",
                        "record_role": "direct_component_log",
                        "defect_flags": "",
                    }
                ],
                "p0_allocation_rows": [],
                "p0_source_metadata": [
                    {
                        "source_id": "source:legacy_component_meter",
                        "source_name": "Legacy component meter",
                        "source_type": "synthetic_factory_energy_record",
                    }
                ],
            },
            ensure_ascii=False,
            sort_keys=True,
            indent=2,
        ),
        encoding="utf-8",
    )
    target_map_path = inputs / "factory-target-map.json"
    target_map_path.write_text(
        json.dumps(
            {
                "schema_version": "1.0",
                "source_ifc_sha256": "LEGACY_IFC_HASH",
                "target_ifc_sha256": actual_extraction.ifc_sha256,
                "mapping_method": "controlled-test-map",
                "component_global_id_map": {legacy_global_id: component.global_id},
            },
            ensure_ascii=False,
            sort_keys=True,
            indent=2,
        ),
        encoding="utf-8",
    )
    config_path = inputs / "mapped-factory-case.json"
    config_path.write_text(
        json.dumps(
            {
                "schemaVersion": "m23-canonical-v2",
                "releaseProfile": "actual-case",
                "releaseReady": False,
                "ifc": _binding(ACTUAL_IFC),
                "factorWorkbook": _binding(factor_path),
                "ontology": _binding(ONTOLOGY),
                "materialEvidence": None,
                "module": {
                    "sourceIdentity": "Main_Modularization_Model",
                    "name": "Main Modularization Model",
                    "identitySource": "case_manifest",
                },
                "requestedScope": "A1-A3",
                "includeOpenings": False,
                "factoryInput": _binding(factory_path),
                "factoryTargetMap": _binding(target_map_path),
                "generatedAtUtc": "2026-07-21T00:00:00Z",
            },
            ensure_ascii=False,
            sort_keys=True,
            indent=2,
        ),
        encoding="utf-8",
    )

    release_dir = build_release(
        load_case_config(config_path),
        tmp_path / "mapped-factory-release",
        "mapped_factory_target",
    )
    graph = json.loads((release_dir / "multigranular_carbon_kg.json").read_text("utf-8"))
    energy = next(
        node
        for node in graph["nodes"]
        if node["props"].get("recordId") == "MAP_REMAP_DIRECT_COMPONENT_001"
        and "EnergyConsumption" in node["labels"]
    )
    product_edges = [
        edge
        for edge in graph["edges"]
        if edge["src"] == energy["id"] and edge["type"] == "recordedForObject"
    ]

    assert energy["props"]["productTargetId"] == component.id
    assert [edge["tgt"] for edge in product_edges] == [component.id]
    assert legacy_global_id not in json.dumps(graph, ensure_ascii=False)
    assert all("FactoryTarget" not in node["labels"] for node in graph["nodes"])
    assert all("Mapping" not in edge["type"] for edge in graph["edges"])


def test_real_a1_material_generator_writes_traceable_schema_v2_evidence(tmp_path: Path):
    from generate_m23_real_a1_material_evidence import generate_material_evidence

    evidence_path = tmp_path / "real-a1-material-evidence.json"
    report_path = tmp_path / "real-a1-material-evidence-report.json"

    result = generate_material_evidence(
        ifc_path=ACTUAL_IFC,
        factor_workbook=ACTUAL_CONFIG.parent
        / ".."
        / ".."
        / "outputs"
        / "research_experiments"
        / "case_study_facts_20260715"
        / "co2e_complete_factors"
        / "Embodied_Carbon_Coefficients_Hong_Kong_CO2e_Complete_2026-07-15.xlsx",
        material_pool=ROOT / "material_pool_enhanced1.xlsx",
        output_path=evidence_path,
        report_path=report_path,
    )

    payload = json.loads(evidence_path.read_text("utf-8"))
    report = json.loads(report_path.read_text("utf-8"))
    records = payload["records"]

    assert payload["schemaVersion"] == "m23-canonical-v2"
    assert result.accepted_record_count > 0
    assert result.rejected_record_count > 0
    assert len(records) == result.accepted_record_count + result.rejected_record_count
    assert report["skippedCounts"]["factor_not_mapped"] > 0
    assert all(record["measured"] is False for record in records)
    assert all(set(record) == release_module._EVIDENCE_RECORD_KEYS for record in records)
    assert any(
        record["formulaCode"] == "volume_density"
        and any(
            operand["role"] == "material_volume"
            and operand["sourceId"] == "ifc:NetVolume"
            and str(operand["designQuantityId"]).startswith("DesignQuantity:")
            for operand in record["operands"]
        )
        and any(
            operand["role"] == "density"
            and str(operand["sourceId"]).startswith("material_pool:")
            for operand in record["operands"]
        )
        for record in records
    )
    assert any(
        record["formulaCode"] == "area_thickness_density"
        and any(operand["role"] == "thickness" for operand in record["operands"])
        for record in records
    )


def _operand_by_role(record: dict[str, object], role: str) -> dict[str, object]:
    matches = [
        operand
        for operand in record["operands"]
        if isinstance(operand, dict) and operand.get("role") == role
    ]
    assert len(matches) == 1
    return matches[0]


def _record_mass_kg(record: dict[str, object]) -> float:
    density = float(_operand_by_role(record, "density")["value"])
    if record["formulaCode"] == "volume_density":
        volume = _operand_by_role(record, "material_volume")
        return float(volume["value"]) * density
    if record["formulaCode"] == "area_thickness_density":
        area = _operand_by_role(record, "material_area")
        thickness = _operand_by_role(record, "thickness")
        thickness_value = float(thickness["value"])
        thickness_unit = str(thickness["unit"])
        thickness_m = thickness_value / 1000.0 if thickness_unit == "mm" else thickness_value
        return float(area["value"]) * thickness_m * density
    raise AssertionError(f"unsupported formula {record['formulaCode']}")


def test_one_layer_named_twice_is_costed_once(tmp_path: Path):
    """A component may carry the same layer under two material associations.

    The model assigns an authoring placeholder material to a whole family of
    walls through one blanket association and then names the actual material on
    each wall separately. Both associations describe one physical layer, and
    each used to be given the full component mass at the same factor, so the
    wall was costed twice. Slabs that genuinely carry two materials, in
    different amounts, must still produce two records.
    """
    from generate_m23_real_a1_material_evidence import (
        _record_mass_kg as record_mass_kg,
        generate_material_evidence,
    )

    evidence_path = tmp_path / "evidence.json"
    report_path = tmp_path / "report.json"
    generate_material_evidence(
        ifc_path=ACTUAL_IFC,
        factor_workbook=ACTUAL_CONFIG.parent
        / ".."
        / ".."
        / "outputs"
        / "research_experiments"
        / "case_study_facts_20260715"
        / "co2e_complete_factors"
        / "Embodied_Carbon_Coefficients_Hong_Kong_CO2e_Complete_2026-07-15.xlsx",
        material_pool=ROOT / "material_pool_enhanced1.xlsx",
        output_path=evidence_path,
        report_path=report_path,
    )
    records = json.loads(evidence_path.read_text("utf-8"))["records"]
    # Records whose mass cannot be computed are the rejected ones; they carry no
    # carbon, so they cannot carry it twice either.
    costed = [
        (record, record_mass_kg(record))
        for record in records
        if record_mass_kg(record) is not None
    ]
    assert len(costed) > 200, "expected the case model to yield a substantial account"

    seen: dict[tuple, str] = {}
    for record, mass in costed:
        signature = (
            record["componentId"],
            record["factorSourceRowId"],
            record["formulaCode"],
            round(mass, 9),
        )
        assert signature not in seen, (
            f"component {record['componentId']} is charged the same mass twice: "
            f"{record['recordId']} repeats {seen[signature]}"
        )
        seen[signature] = record["recordId"]

    report = json.loads(report_path.read_text("utf-8"))
    dropped = [
        row
        for row in report["skippedRecords"]
        if row.get("reasonCode") == "duplicate_material_layer"
    ]
    assert dropped, "the case model does contain repeated layers, so some must be dropped"
    assert all(row.get("retainedAssociationId") for row in dropped)

    # The composite slab is the case that must survive: concrete and profiled
    # steel sheet on one component, in genuinely different amounts.
    by_component: dict[str, list[float]] = {}
    for record, mass in costed:
        by_component.setdefault(record["componentId"], []).append(round(mass, 6))
    composite = [
        masses for masses in by_component.values() if len(set(masses)) > 1
    ]
    assert composite, "slabs carrying two materials in different amounts must be kept"


def test_real_a1_generator_corrects_g1_slab_profiled_sheet_to_concrete_and_thin_formwork(
    tmp_path: Path,
):
    from dm2c_m23_ifc import CanonicalIFCExtractor
    from generate_m23_real_a1_material_evidence import generate_material_evidence

    evidence_path = tmp_path / "real-a1-material-evidence.json"
    report_path = tmp_path / "real-a1-material-evidence-report.json"

    generate_material_evidence(
        ifc_path=ACTUAL_IFC,
        factor_workbook=ACTUAL_CONFIG.parent
        / ".."
        / ".."
        / "outputs"
        / "research_experiments"
        / "case_study_facts_20260715"
        / "co2e_complete_factors"
        / "Embodied_Carbon_Coefficients_Hong_Kong_CO2e_Complete_2026-07-15.xlsx",
        material_pool=ROOT / "material_pool_enhanced1.xlsx",
        output_path=evidence_path,
        report_path=report_path,
    )

    extraction = CanonicalIFCExtractor(ACTUAL_IFC).extract()
    g1_slab = next(
        component
        for component in extraction.components
        if component.ifc_class == "IfcSlab" and "2135000" in component.name
    )
    records = json.loads(evidence_path.read_text("utf-8"))["records"]
    report = json.loads(report_path.read_text("utf-8"))
    slab_records = [record for record in records if record["componentId"] == g1_slab.id]

    assert not any(
        record["formulaCode"] == "volume_density"
        and _operand_by_role(record, "density")["value"] == 7850
        and _operand_by_role(record, "material_volume")["value"] == pytest.approx(
            2.694096406688321
        )
        for record in slab_records
    )

    concrete = [
        record
        for record in slab_records
        if record["formulaCode"] == "volume_density"
        and _operand_by_role(record, "density")["value"] == 2500
    ]
    thin_steel = [
        record
        for record in slab_records
        if record["formulaCode"] == "area_thickness_density"
        and _operand_by_role(record, "density")["value"] == 7850
        and _operand_by_role(record, "thickness")["value"] == pytest.approx(1.0)
        and _operand_by_role(record, "thickness")["unit"] == "mm"
    ]

    assert len(concrete) == 1
    assert _operand_by_role(concrete[0], "material_volume")["sourceId"] == "ifc:NetVolume"
    assert _record_mass_kg(concrete[0]) == pytest.approx(2.694096406688321 * 2500)
    assert len(thin_steel) == 1
    assert _operand_by_role(thin_steel[0], "material_area")["sourceId"] == "ifc:NetArea"
    assert _record_mass_kg(thin_steel[0]) == pytest.approx(
        24.043491190339882 * 0.001 * 7850
    )
    assert report["suspiciousRecords"] == []
    assert report["notes"]["slabReinforcement"] == (
        "slab reinforcement excluded; Pset_ReinforcementBarPitchOfSlab can support future calculation"
    )
    assert report["materialBreakdown"]["Concrete"]["acceptedRecords"] >= 1
    # The band brackets the steel a module of this size plausibly holds. It was
    # widened downwards once duplicate material layers stopped being counted: a
    # component carrying the same steel under two material associations used to
    # contribute its mass twice.
    assert report["materialBreakdown"]["Steel"]["massKg"] == pytest.approx(
        2350,
        abs=450,
    )


def test_real_a1_material_and_synthetic_a3_fixture_release_is_publishable(
    tmp_path: Path,
):
    from generate_m23_real_a1_material_evidence import (
        generate_material_evidence,
        write_controlled_fixture_case,
    )

    evidence_path = tmp_path / "real-a1-material-evidence.json"
    report_path = tmp_path / "real-a1-material-evidence-report.json"
    generate_material_evidence(
        ifc_path=ACTUAL_IFC,
        factor_workbook=ACTUAL_CONFIG.parent
        / ".."
        / ".."
        / "outputs"
        / "research_experiments"
        / "case_study_facts_20260715"
        / "co2e_complete_factors"
        / "Embodied_Carbon_Coefficients_Hong_Kong_CO2e_Complete_2026-07-15.xlsx",
        material_pool=ROOT / "material_pool_enhanced1.xlsx",
        output_path=evidence_path,
        report_path=report_path,
    )
    fixture_config_path = tmp_path / "m23_canonical_v2_case_fixture.json"
    write_controlled_fixture_case(
        base_case_config=ACTUAL_CONFIG,
        material_evidence=evidence_path,
        output_path=fixture_config_path,
        generated_at_utc="2026-07-22T00:00:00Z",
    )

    release_dir = build_release(
        load_case_config(fixture_config_path),
        tmp_path / "fixture-release",
        "real_a1_material_fixture",
    )
    graph = json.loads((release_dir / "multigranular_carbon_kg.json").read_text("utf-8"))
    stats = json.loads(
        (release_dir / "multigranular_carbon_kg_stats.json").read_text("utf-8")
    )
    manifest = _read_manifest(release_dir)
    validation_rows = list(
        csv.DictReader(
            (release_dir / "multigranular_carbon_kg_validation.csv").open(
                encoding="utf-8"
            )
        )
    )

    accepted = stats["validation"]["acceptedByKind"]
    rejected = stats["validation"]["rejectedByKind"]
    assert manifest["releaseProfile"] == "controlled-fixture"
    assert manifest["releaseReady"] is True
    assert manifest["syntheticFactoryInputsUsed"] is True
    assert manifest["gates"]["alignment"] == "pass"
    assert accepted["material"] > 0
    assert accepted["energy"] == 17
    assert rejected["material"] > 0
    assert stats["applicationClassCounts"]["MaterialConsumption"] == accepted["material"]
    assert stats["applicationClassCounts"]["EnergyConsumption"] == 17
    assert stats["applicationClassCounts"]["CarbonEmission"] == (
        accepted["material"] + 17
    )
    assert stats["relationCounts"]["ofMaterial"] >= accepted["material"]
    assert stats["relationCounts"]["hasFactor"] >= accepted["material"] + 17
    assert stats["calculationCounts"]["material"] == accepted["material"]
    assert any(
        row["kind"] == "material"
        and row["status"] == "accepted"
        and '"measured":false' in row["evidenceJson"]
        and "ifc:NetVolume" in row["evidenceJson"]
        and "material_pool:" in row["evidenceJson"]
        for row in validation_rows
    )

    nodes = {node["id"]: node for node in graph["nodes"]}
    material_total = 0.0
    energy_total = 0.0
    for emission in (
        node for node in graph["nodes"] if "CarbonEmission" in node["labels"]
    ):
        drivers = [
            edge["tgt"]
            for edge in graph["edges"]
            if edge["src"] == emission["id"] and edge["type"] == "hasCarbonDriver"
        ]
        assert len(drivers) == 1
        driver = nodes[drivers[0]]
        if "MaterialConsumption" in driver["labels"]:
            material_total += emission["props"]["emissionValue"]
        elif "EnergyConsumption" in driver["labels"]:
            energy_total += emission["props"]["emissionValue"]
    # Lowered from 8,000 when duplicate material layers stopped being counted.
    assert 7_000 < material_total < 12_000
    assert energy_total > 0


def test_case_config_rejects_unknown_keys_and_input_hash_mismatch(
    tmp_path: Path, controlled_case
):
    payload = dict(controlled_case["configPayload"])
    payload["legacyFactoryPath"] = "p0_factory_synth.xlsx"
    unknown_path = tmp_path / "unknown-key.json"
    unknown_path.write_text(json.dumps(payload), encoding="utf-8")

    with pytest.raises(CaseConfigError, match="unknown"):
        load_case_config(unknown_path)

    mismatch = json.loads(json.dumps(controlled_case["configPayload"]))
    mismatch["ifc"]["sha256"] = "0" * 64
    mismatch_path = tmp_path / "hash-mismatch.json"
    mismatch_path.write_text(json.dumps(mismatch), encoding="utf-8")
    out_root = tmp_path / "hash-mismatch-output"

    with pytest.raises(ReleaseBuildError) as captured:
        build_release(load_case_config(mismatch_path), out_root, "hash_mismatch")

    assert captured.value.reason_code == "input_hash_mismatch"
    assert not (out_root / "hash_mismatch").exists()
    _assert_no_staging_dirs(out_root, "hash_mismatch")


def test_control_json_rejects_duplicate_keys_and_nonfinite_constants(
    tmp_path: Path, controlled_case, actual_extraction
):
    payload = controlled_case["configPayload"]
    duplicate_config = tmp_path / "duplicate-config.json"
    duplicate_text = json.dumps(payload, ensure_ascii=False, indent=2).replace(
        '"schemaVersion": "m23-canonical-v2",',
        '"schemaVersion": "m23-canonical-v2",\n  "schemaVersion": "duplicate",',
        1,
    )
    duplicate_config.write_text(duplicate_text, encoding="utf-8")
    with pytest.raises(CaseConfigError, match="duplicate"):
        load_case_config(duplicate_config)

    nonfinite_payload = json.loads(json.dumps(payload))
    nonfinite_payload["generatedAtUtc"] = math.nan
    nonfinite_config = tmp_path / "nonfinite-config.json"
    nonfinite_config.write_text(json.dumps(nonfinite_payload), encoding="utf-8")
    with pytest.raises(CaseConfigError, match="non-finite"):
        load_case_config(nonfinite_config)

    evidence_path = tmp_path / "duplicate-evidence.json"
    duplicate_evidence = controlled_case["evidencePath"].read_text("utf-8").replace(
        '"records": [', '"records": [],\n  "records": [', 1
    )
    evidence_path.write_text(duplicate_evidence, encoding="utf-8")
    evidence_config_payload = json.loads(json.dumps(payload))
    evidence_config_payload["materialEvidence"] = _binding(evidence_path)
    evidence_config = tmp_path / "duplicate-evidence-config.json"
    evidence_config.write_text(json.dumps(evidence_config_payload), encoding="utf-8")
    with pytest.raises(ReleaseBuildError) as captured:
        assemble_material_evidence(
            load_case_config(evidence_config),
            actual_extraction,
            FactorLibrary(controlled_case["factorPath"]),
        )
    assert captured.value.reason_code == "material_evidence_schema_invalid"


@pytest.mark.parametrize(
    ("case_name", "operand_index", "field", "invalid_value"),
    (
        ("numeric-string-value", 1, "value", "2700.0"),
        ("numeric-operand-id", 1, "operandId", 123),
        ("numeric-source-id", 1, "sourceId", 456),
        ("wrong-design-quantity-id-type", 0, "designQuantityId", 789),
        ("unclassified-numeric-value", 1, "value", {"not": "numeric"}),
    ),
)
def test_material_evidence_operand_schema_rejects_coercible_and_invalid_json_types(
    tmp_path: Path,
    controlled_case,
    actual_extraction,
    case_name: str,
    operand_index: int,
    field: str,
    invalid_value: object,
):
    record = json.loads(json.dumps(controlled_case["validRecord"]))
    record["operands"][operand_index][field] = invalid_value
    evidence_path = tmp_path / f"{case_name}.json"
    evidence_path.write_text(
        json.dumps(
            {"schemaVersion": "m23-canonical-v2", "records": [record]},
            ensure_ascii=False,
            sort_keys=True,
            indent=2,
        ),
        encoding="utf-8",
    )
    config_payload = json.loads(json.dumps(controlled_case["configPayload"]))
    config_payload["materialEvidence"] = _binding(evidence_path)
    config_path = tmp_path / f"{case_name}-config.json"
    config_path.write_text(
        json.dumps(config_payload, ensure_ascii=False, sort_keys=True, indent=2),
        encoding="utf-8",
    )

    with pytest.raises(ReleaseBuildError) as captured:
        assemble_material_evidence(
            load_case_config(config_path),
            actual_extraction,
            FactorLibrary(controlled_case["factorPath"]),
        )

    assert captured.value.reason_code == "material_evidence_schema_invalid"


def test_evidence_assembly_resolves_exact_ifc_ownership_and_factor_source_row(
    controlled_case, actual_extraction
):
    config = load_case_config(controlled_case["configPath"])
    factors = FactorLibrary(controlled_case["factorPath"])

    assembly = assemble_material_evidence(config, actual_extraction, factors)

    assert isinstance(assembly, EvidenceAssembly)
    assert len(assembly.candidates) == 1
    assert len(assembly.rejected) == 1
    assert assembly.source_ledger_identities == (
        "controlled-material:accepted",
        "controlled-material:rejected-ownership",
    )
    candidate = assembly.candidates[0]
    assert candidate.source_record_id == controlled_case["association"].id
    assert candidate.product_target_id == controlled_case["association"].component_id
    assert candidate.material_id == controlled_case["association"].material_id
    assert candidate.factor.source_row_id == controlled_case["factorSourceRowId"]
    assert candidate.operands[0].design_quantity_id == controlled_case["quantity"].id
    assert assembly.rejected[0].reason_code == "evidence_ownership_mismatch"

    calculated = validate_candidates(assembly.candidates)
    rows = build_validation_rows(
        ValidationResultSet(
            accepted=calculated.accepted,
            rejected=assembly.rejected + calculated.rejected,
        )
    )
    rejected_row = next(row for row in rows if row["status"] == "rejected")
    evidence_source_id = f"sha256:{config.material_evidence.sha256}"
    assert rejected_row["sourceRecordId"] == controlled_case["rejectedAssociation"].id
    assert rejected_row["evidenceSourceId"] == evidence_source_id
    assert rejected_row["sourceIdentity"]
    with pytest.raises(FrozenInstanceError):
        assembly.candidates = ()  # type: ignore[misc]


def test_calculation_rejection_derives_stable_source_identity_in_row_and_evidence():
    evidence_source_id = "sha256:" + "C" * 64
    source_record_id = "MaterialAssociation:calculation-rejection"
    rows = build_validation_rows(
        ValidationResultSet(
            accepted=(),
            rejected=(
                ValidationIssue(
                    record_id="rejected:formula",
                    reason_code="formula_unsupported",
                    message="unsupported formula",
                    evidence={
                        "kind": "material",
                        "sourceRecordId": source_record_id,
                        "evidenceSourceId": evidence_source_id,
                        "formulaCode": "unsupported_formula",
                    },
                ),
            ),
        )
    )
    expected = stable_id("SourceRecord", evidence_source_id, source_record_id)

    assert rows[0]["sourceIdentity"] == expected
    assert json.loads(rows[0]["evidenceJson"])["sourceIdentity"] == expected


def test_controlled_release_reconciles_validation_ledger_and_rejected_exclusion(
    tmp_path: Path, controlled_case
):
    release_dir = build_release(
        load_case_config(controlled_case["configPath"]),
        tmp_path / "releases",
        "controlled_reconciliation",
    )

    with (release_dir / "multigranular_carbon_kg_validation.csv").open(
        encoding="utf-8", newline=""
    ) as handle:
        rows = list(csv.DictReader(handle))
    graph = json.loads((release_dir / "multigranular_carbon_kg.json").read_text("utf-8"))
    manifest = _read_manifest(release_dir)

    assert [row["recordId"] for row in rows] == [
        "controlled-material:accepted",
        "controlled-material:rejected-ownership",
    ]
    assert [row["status"] for row in rows] == ["accepted", "rejected"]
    assert rows[0]["reasonCode"] == ""
    assert rows[0]["consumptionId"]
    assert rows[0]["quantityId"]
    assert rows[0]["factorId"]
    assert rows[0]["emissionId"]
    assert rows[1]["reasonCode"] == "evidence_ownership_mismatch"
    assert rows[1]["consumptionId"] == ""
    assert rows[1]["emissionId"] == ""
    assert json.loads(rows[0]["evidenceJson"])["formulaCode"] == "volume_density"
    assert json.loads(rows[1]["evidenceJson"])["associationId"]
    serialized_graph = json.dumps(graph, ensure_ascii=False)
    assert "controlled-material:accepted" in serialized_graph
    assert "controlled-material:rejected-ownership" not in serialized_graph
    assert sum("MaterialConsumption" in node["labels"] for node in graph["nodes"]) == 1
    assert manifest["coverage"] == {
        "acceptedByKind": {"material": 1},
        "acceptedCount": 1,
        "candidateCount": 2,
        "coverageFormula": "acceptedCount / candidateCount",
        "coverageValue": 0.5,
        "rejectedByKind": {"material": 1},
        "rejectedByReason": {"evidence_ownership_mismatch": 1},
        "rejectedCount": 1,
    }


def test_release_is_byte_deterministic_across_output_roots_and_manifest_is_exact(
    tmp_path: Path, controlled_case
):
    config = load_case_config(controlled_case["configPath"])
    first = build_release(config, tmp_path / "root-a", "deterministic_release")
    second = build_release(config, tmp_path / "root-b", "deterministic_release")

    assert {path.name for path in first.iterdir()} == ARTIFACT_NAMES
    assert {path.name for path in second.iterdir()} == ARTIFACT_NAMES
    for filename in sorted(ARTIFACT_NAMES):
        assert (first / filename).read_bytes() == (second / filename).read_bytes()

    manifest = _read_manifest(first)
    stats = json.loads((first / "multigranular_carbon_kg_stats.json").read_text("utf-8"))
    assert manifest["schemaVersion"] == "m23-canonical-v2"
    assert manifest["releaseProfile"] == "controlled-fixture"
    assert manifest["releaseReady"] is True
    assert manifest["syntheticFactoryInputsUsed"] is False
    assert set(manifest["outputs"]) == set(OUTPUT_NAMES)
    assert set(manifest["gates"]) == {
        "alignment",
        "structuralCypher",
        "liveCypherRoundTrip",
    }
    assert manifest["gates"] == {
        "alignment": "pass",
        "structuralCypher": "pass",
        "liveCypherRoundTrip": "not_configured",
    }
    assert "manifest" not in manifest["outputs"]
    assert "case_version_manifest.json" not in json.dumps(manifest["outputs"])
    for key, filename in OUTPUT_NAMES.items():
        descriptor = manifest["outputs"][key]
        assert descriptor["path"] == filename
        assert Path(descriptor["path"]).name == descriptor["path"]
        assert descriptor["sha256"] == _hash_size(first / filename)[0]
        assert descriptor["sizeBytes"] == _hash_size(first / filename)[1]
        assert descriptor["sha256"] == descriptor["sha256"].upper()
    for descriptor in manifest["inputs"].values():
        if descriptor is not None:
            assert set(descriptor) >= {"sha256", "sizeBytes"}
            assert descriptor["sha256"] == descriptor["sha256"].upper()
    assert set(stats["applicationClassCounts"]) == set(APPLICATION_CLASSES)
    assert set(stats["relationCounts"]) == set(PRINCIPAL_PREDICATES) | set(
        OPTIONAL_CONTEXT_PREDICATES
    )


def test_nested_provenance_is_canonical_json_text_and_property_gate_is_fail_closed(
    tmp_path: Path, controlled_case
):
    graph = CanonicalLPGGraph()
    with pytest.raises(TypeError, match="Neo4j"):
        graph.add_node("bad:mapping", ["BuildingComponent"], {"nested": {"x": 1}})
    with pytest.raises(TypeError, match="Neo4j"):
        graph.add_node("bad:list", ["BuildingComponent"], {"nested": [{"x": 1}]})
    with pytest.raises(TypeError, match="homogeneous"):
        graph.add_node("bad:mixed", ["BuildingComponent"], {"mixed": [1, "2"]})
    with pytest.raises(ValueError, match="finite"):
        graph.add_node("bad:nan", ["BuildingComponent"], {"value": math.nan})
    with pytest.raises(ValueError, match="64-bit"):
        graph.add_node("bad:integer", ["BuildingComponent"], {"value": 2**63})
    with pytest.raises(TypeError, match="property key"):
        graph.add_node("bad:key", ["BuildingComponent"], {1: "not-a-string"})

    release_dir = build_release(
        load_case_config(controlled_case["configPath"]),
        tmp_path / "provenance-release",
        "provenance_contract",
    )
    payload = json.loads((release_dir / "multigranular_carbon_kg.json").read_text("utf-8"))
    cypher = (release_dir / "multigranular_carbon_kg.cypher").read_text("utf-8")
    quantity = next(node for node in payload["nodes"] if "ConsumptionQuantity" in node["labels"])

    for field in ("orderedRawOperands", "conversionSteps"):
        values = quantity["props"][field]
        assert isinstance(values, list) and values
        assert all(isinstance(value, str) for value in values)
        decoded = [json.loads(value) for value in values]
        assert all(isinstance(value, dict) for value in decoded)
        assert values == [
            json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
            for value in decoded
        ]
        assert all(value.replace("\\", "\\\\").replace("'", "\\'") in cypher for value in values)


@pytest.mark.parametrize(
    "release_id",
    (
        "../escape",
        "nested/name",
        r"nested\name",
        "C:drive",
        ".",
        "..",
        "CON",
        "CON.txt",
        "trailing.",
    ),
)
def test_invalid_release_id_is_rejected_without_publication(
    tmp_path: Path, controlled_case, release_id: str
):
    out_root = tmp_path / "unsafe-releases"
    with pytest.raises(ReleaseBuildError) as captured:
        build_release(load_case_config(controlled_case["configPath"]), out_root, release_id)

    assert captured.value.reason_code == "invalid_release_id"
    assert not out_root.exists() or not list(out_root.iterdir())


def test_existing_and_concurrent_destination_are_never_overwritten(
    tmp_path: Path, controlled_case, monkeypatch: pytest.MonkeyPatch
):
    config = load_case_config(controlled_case["configPath"])
    out_root = tmp_path / "no-overwrite"
    existing = out_root / "already_exists"
    existing.mkdir(parents=True)
    sentinel = existing / "sentinel.txt"
    sentinel.write_text("preserve", encoding="utf-8")

    with pytest.raises(ReleaseBuildError) as captured:
        build_release(config, out_root, "already_exists")
    assert captured.value.reason_code == "destination_exists"
    assert sentinel.read_text("utf-8") == "preserve"

    real_audit = release_module.audit_release
    concurrent_final = out_root / "concurrent_destination"

    def concurrent_audit(release_dir: Path, *, live_cypher_config=None):
        report = real_audit(release_dir, live_cypher_config=live_cypher_config)
        concurrent_final.mkdir()
        (concurrent_final / "sentinel.txt").write_text("concurrent", encoding="utf-8")
        return report

    monkeypatch.setattr(release_module, "audit_release", concurrent_audit)
    with pytest.raises(ReleaseBuildError) as concurrent_error:
        build_release(config, out_root, "concurrent_destination")
    assert concurrent_error.value.reason_code == "destination_exists"
    assert (concurrent_final / "sentinel.txt").read_text("utf-8") == "concurrent"
    _assert_no_staging_dirs(out_root, "concurrent_destination")


def test_late_alignment_failure_removes_staging_and_never_creates_final(
    tmp_path: Path, controlled_case, monkeypatch: pytest.MonkeyPatch
):
    config = load_case_config(controlled_case["configPath"])
    out_root = tmp_path / "late-gate"
    real_audit = release_module.audit_release

    def failed_audit(release_dir: Path, *, live_cypher_config=None):
        report = real_audit(release_dir, live_cypher_config=live_cypher_config)
        return {
            **report,
            "highSeverityViolationCount": 1,
            "findings": [
                {
                    "severity": "high",
                    "code": "injected_late_failure",
                    "artifact": "release",
                    "entityId": "",
                    "message": "controlled late gate failure",
                    "evidence": {},
                }
            ],
        }

    monkeypatch.setattr(release_module, "audit_release", failed_audit)
    with pytest.raises(ReleaseBuildError) as captured:
        build_release(config, out_root, "late_failure")

    assert captured.value.reason_code == "alignment_gate_failed"
    assert not (out_root / "late_failure").exists()
    _assert_no_staging_dirs(out_root, "late_failure")


def test_release_records_structural_cypher_and_no_live_gate_claim(
    tmp_path: Path, controlled_case
):
    release_dir = build_release(
        load_case_config(controlled_case["configPath"]),
        tmp_path / "cypher-gates",
        "cypher_gate_contract",
    )
    manifest = _read_manifest(release_dir)
    report = json.loads((release_dir / "m2_alignment_report.json").read_text("utf-8"))
    cypher = (release_dir / "multigranular_carbon_kg.cypher").read_text("utf-8")
    graph = json.loads((release_dir / "multigranular_carbon_kg.json").read_text("utf-8"))

    assert manifest["gates"]["structuralCypher"] == "pass"
    assert manifest["gates"]["liveCypherRoundTrip"] == "not_configured"
    assert report["gates"]["structuralCypher"] == "pass"
    assert report["gates"]["liveCypherRoundTrip"] == "not_configured"
    assert "DETACH DELETE" not in cypher
    assert all(edge["occurrenceId"] in cypher for edge in graph["edges"])


def test_release_records_are_immutable_and_legacy_cli_delegates_locally(
    controlled_case, monkeypatch: pytest.MonkeyPatch
):
    config = load_case_config(controlled_case["configPath"])
    assert isinstance(config, CaseConfig)
    artifacts = ReleaseArtifacts(
        graph_json=Path("graph.json"),
        cypher=Path("graph.cypher"),
        stats=Path("stats.json"),
        validation=Path("validation.csv"),
        alignment_report=Path("alignment.json"),
        manifest=Path("manifest.json"),
    )
    with pytest.raises(FrozenInstanceError):
        artifacts.graph_json = Path("other.json")  # type: ignore[misc]

    import dm2c_multigranular_carbon_kg as legacy_entry

    calls: list[object] = []

    def fake_release_main(argv=None):
        calls.append(argv)
        return 17

    monkeypatch.setattr(release_module, "main", fake_release_main)
    assert legacy_entry.main(["--case-config", str(controlled_case["configPath"])]) == 17
    assert calls == [["--case-config", str(controlled_case["configPath"])]]
    assert not hasattr(legacy_entry, "parse_args")
