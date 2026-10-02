#!/usr/bin/env python3
"""
DM2C multi-granularity carbon knowledge graph builder.

This script is research-specific for M2.3. It keeps the IFC-derived design
backbone separate from the carbon calculation layer:

1. IFC design backbone:
   ModularUnit -> BuildingComponent -> Material / DesignQuantity

2. Optional canonical manufacturing context:
   ManufacturingProcessTemplate -> ProductionStage -> ManufacturingActivity
   -> ManufacturingResource

3. Source-fact calculation layer:
   MaterialConsumption / EnergyConsumption -> ConsumptionQuantity +
   EmissionFactor, with CarbonEmission -[:hasCarbonDriver]-> Consumption.

Carbon values are represented as traceable source facts. Product and process
totals are query-time projections and are never materialized as aggregate nodes.
"""

from __future__ import annotations

from copy import deepcopy
import csv
import hashlib
import json
import math
import re
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Set, Tuple

from dm2c_m23_ifc import (
    CanonicalIFCExtractor,
    ComponentRecord as CanonicalComponentRecord,
    ComponentTypeRecord,
    DesignQuantityRecord,
    IFCExtractionResult,
    MaterialAssociationRecord,
    deduplicate_material_associations,
)
from dm2c_m23_units import UnitError, parse_factor_denominator, validate_quantity_factor
from dm2c_m23_calculation import (
    AcceptedConsumption,
    CandidateValidation,
    ConsumptionCandidate,
    QuantityOperand,
    ValidationIssue,
    ValidationResultSet,
    validate_candidates,
    validate_energy_candidate,
    validate_material_candidate,
)
from dm2c_m23_canonical import CanonicalLPGGraph, stable_id
from dm2c_m23_accounting import (
    EnergyAllocationTarget,
    EnergyPopulationPlan,
    validate_energy_population_plan,
)

try:
    import ifcopenshell
except ImportError:  # pragma: no cover - reported at runtime
    ifcopenshell = None

try:
    import openpyxl
except ImportError:  # pragma: no cover - factor xlsx becomes optional
    openpyxl = None


DEFAULT_ELEMENT_TYPES = [
    "IfcWall",
    "IfcWallStandardCase",
    "IfcBeam",
    "IfcColumn",
    "IfcSlab",
    "IfcDoor",
    "IfcWindow",
    "IfcBuildingElementProxy",
    "IfcCableCarrierSegment",
    "IfcCableCarrierFitting",
    "IfcPipeSegment",
    "IfcPipeFitting",
    "IfcLightFixture",
]


EXCLUDED_FACTORY_BACKBONE_TYPES = {
    "IfcFurniture",
    "IfcAlarm",
    "IfcFireSuppressionTerminal",
    "IfcAirTerminal",
    "IfcValve",
}


PLACEHOLDER_MATERIALS = {
    "",
    "unnamed",
    "none",
    "null",
    "n/a",
    "na",
    "material",
    "ifcmaterial",
    "default",
    "x",
    "未命名",
    "未定义",
    "未指定",
    "默认",
    "默认材料",
    "材料",
    "无",
}


MATERIAL_KEYWORDS = [
    ("steel", ["steel", "stainless", "rebar", "metal", "金属", "钢", "鋼", "铁", "鐵"], 7850.0),
    ("aluminium", ["aluminium", "aluminum", "alum", "铝", "鋁"], 2700.0),
    ("concrete", ["concrete", "混凝土", "砼"], 2400.0),
    ("cement", ["cement", "水泥"], 1450.0),
    ("mortar", ["mortar", "砂浆", "砂漿"], 1800.0),
    ("glass", ["glass", "玻璃"], 2500.0),
    ("plastic", ["plastic", "pvc", "vinyl", "resin", "塑料", "乙烯", "树脂", "樹脂"], 1200.0),
    ("gypsum", ["gypsum", "plasterboard", "石膏"], 950.0),
    ("ceramics", ["ceramic", "tile", "tiles", "瓷", "砖", "磚"], 2000.0),
    ("timber", ["timber", "wood", "木"], 600.0),
    ("insulation", ["insulation", "rock wool", "glass wool", "保温", "保溫", "岩棉"], 80.0),
]


def normalize_text(text: Optional[str]) -> str:
    return re.sub(r"\s+", "", str(text or "").strip().casefold())


def safe_float(value: Any, default: float = 0.0) -> float:
    try:
        if value is None or value == "":
            return default
        return float(value)
    except Exception:
        return default


def slugify(text: Any, max_len: int = 96) -> str:
    token = re.sub(r"[^A-Za-z0-9_]+", "_", str(text)).strip("_")
    return (token or "x")[:max_len]


def sanitize_key(text: Any) -> str:
    key = re.sub(r"[^A-Za-z0-9_]+", "_", str(text)).strip("_")
    if not key:
        key = "value"
    if key[0].isdigit():
        key = f"p_{key}"
    return key


def sanitize_label(text: Any) -> str:
    label = re.sub(r"[^A-Za-z0-9_]+", "_", str(text)).strip("_")
    if not label:
        label = "Entity"
    if label[0].isdigit():
        label = f"L_{label}"
    return label


def sanitize_rel(text: Any) -> str:
    rel = re.sub(r"[^A-Za-z0-9_]+", "_", str(text)).strip("_").upper()
    if not rel:
        rel = "RELATED_TO"
    if rel[0].isdigit():
        rel = f"R_{rel}"
    return rel


def cypher_key(key: str) -> str:
    if re.match(r"^[A-Za-z_][A-Za-z0-9_]*$", key):
        return key
    return f"`{key.replace('`', '``')}`"


def cypher_value(value: Any) -> str:
    if value is None:
        return "null"
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return repr(value) if math.isfinite(value) else "null"
    if isinstance(value, list):
        return "[" + ", ".join(cypher_value(v) for v in value) + "]"
    return json.dumps(str(value), ensure_ascii=False)


def first_non_empty(values: Iterable[Any], default: str = "") -> str:
    for value in values:
        text = str(value or "").strip()
        if text:
            return text
    return default


def is_placeholder_material(text: Optional[str]) -> bool:
    raw = str(text or "").strip()
    norm = normalize_text(raw).replace("<", "").replace(">", "").replace("_", "")
    return norm in PLACEHOLDER_MATERIALS


def keyword_family(text: str) -> Tuple[str, float]:
    norm = normalize_text(text)
    for family, keywords, density in MATERIAL_KEYWORDS:
        if any(normalize_text(k) in norm for k in keywords):
            return family, density
    return "unknown", 1200.0


@dataclass
class MaterialRecord:
    name: str
    source: str
    thickness: Optional[float] = None
    status: str = "resolved"


@dataclass
class ComponentRecord:
    global_id: str
    ifc_type: str
    name: str
    object_type: str
    predefined_type: str
    tag: str
    storey: str
    group_names: List[str]
    materials: List[MaterialRecord] = field(default_factory=list)
    quantities: Dict[str, float] = field(default_factory=dict)


@dataclass
class FactorRecord:
    factor_id: str
    keyword: str
    factor_value: float
    factor_unit: str
    source: str
    match_status: str
    confidence: float
    source_row_id: str = ""
    original_factor_unit: str = ""
    denominator_unit: str = ""
    normalized_factor_value: Optional[float] = None
    normalized_denominator: str = ""
    geography: str = ""
    year: str = ""
    system_boundary: str = ""
    source_reference: str = ""
    proxy_status: str = ""
    validation_status: str = ""


class LPGGraph:
    def __init__(self) -> None:
        self.nodes: Dict[str, Dict[str, Any]] = {}
        self.edges: List[Dict[str, Any]] = []
        self._edge_seen: Set[Tuple[str, str, str, str]] = set()

    def add_node(self, node_id: str, labels: List[str], props: Optional[Dict[str, Any]] = None) -> None:
        clean_labels = sorted({sanitize_label(x) for x in labels if str(x or "").strip()})
        clean_props: Dict[str, Any] = {}
        for key, value in (props or {}).items():
            if value is None or value == "":
                continue
            clean_props[sanitize_key(key)] = value

        if node_id not in self.nodes:
            self.nodes[node_id] = {"id": node_id, "labels": clean_labels, "props": clean_props}
            return

        self.nodes[node_id]["labels"] = sorted(set(self.nodes[node_id]["labels"]) | set(clean_labels))
        self.nodes[node_id]["props"].update(clean_props)

    def add_edge(self, src: str, rel_type: str, tgt: str, props: Optional[Dict[str, Any]] = None) -> None:
        rel = sanitize_rel(rel_type)
        clean_props: Dict[str, Any] = {}
        for key, value in (props or {}).items():
            if value is None or value == "":
                continue
            clean_props[sanitize_key(key)] = value
        edge_key = (src, rel, tgt, json.dumps(clean_props, sort_keys=True, ensure_ascii=False))
        if edge_key in self._edge_seen:
            return
        self._edge_seen.add(edge_key)
        self.edges.append(
            {"id": f"e_{len(self.edges) + 1}", "src": src, "type": rel, "tgt": tgt, "props": clean_props}
        )

    def export_json(self, path: Path) -> None:
        path.write_text(
            json.dumps({"nodes": list(self.nodes.values()), "edges": self.edges}, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

    def export_cypher(self, path: Path, wipe_before_import: bool = False) -> None:
        lines: List[str] = [
            "// DM2C multi-granularity carbon knowledge graph",
            "// Generated by dm2c_multigranular_carbon_kg.py",
            "",
        ]
        if wipe_before_import:
            lines.extend(["MATCH (n) DETACH DELETE n;", ""])
        lines.extend(["CREATE CONSTRAINT IF NOT EXISTS FOR (n:DM2CEntity) REQUIRE n.id IS UNIQUE;", ""])

        lines.append("// Nodes")
        for node_id in sorted(self.nodes):
            node = self.nodes[node_id]
            labels = ":".join(["DM2CEntity"] + sorted(node.get("labels", [])))
            lines.append(f"MERGE (n:{labels} {{id: {cypher_value(node_id)}}})")
            props = node.get("props", {})
            if props:
                prop_expr = "{ " + ", ".join(
                    f"{cypher_key(k)}: {cypher_value(v)}" for k, v in sorted(props.items())
                ) + " }"
                lines.append(f"SET n += {prop_expr};")
            else:
                lines.append(";")
        lines.append("")

        lines.append("// Relationships")
        for edge in self.edges:
            lines.append(
                f"MATCH (a:DM2CEntity {{id: {cypher_value(edge['src'])}}}), "
                f"(b:DM2CEntity {{id: {cypher_value(edge['tgt'])}}})"
            )
            lines.append(f"MERGE (a)-[r:{sanitize_rel(edge['type'])}]->(b)")
            props = edge.get("props", {})
            if props:
                prop_expr = "{ " + ", ".join(
                    f"{cypher_key(k)}: {cypher_value(v)}" for k, v in sorted(props.items())
                ) + " }"
                lines.append(f"SET r += {prop_expr};")
            else:
                lines.append(";")
        path.write_text("\n".join(lines), encoding="utf-8")

    def label_counts(self) -> Dict[str, int]:
        counts = Counter()
        for node in self.nodes.values():
            for label in node.get("labels", []):
                counts[label] += 1
        return dict(counts.most_common())

    def edge_counts(self) -> Dict[str, int]:
        return dict(Counter(edge["type"] for edge in self.edges).most_common())


class FactorLibrary:
    def __init__(self, factor_xlsx: Optional[Path]) -> None:
        self.material_factors: List[Dict[str, Any]] = []
        self.energy_factors: List[Dict[str, Any]] = []
        self.load_error = ""
        self.workbook_sha256 = ""
        self.last_strict_resolution_reason = "not_resolved"
        if factor_xlsx:
            self._load_xlsx(factor_xlsx)

    def _load_xlsx(self, path: Path) -> None:
        if openpyxl is None:
            self.load_error = "openpyxl is not installed"
            return
        if not path.exists():
            self.load_error = f"factor file not found: {path}"
            return
        self.workbook_sha256 = hashlib.sha256(path.read_bytes()).hexdigest().upper()
        try:
            wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        except Exception as exc:
            self.load_error = f"failed to load factor file: {exc}"
            return

        if "Embodied Carbon Coefficients" in wb.sheetnames:
            rows = list(wb["Embodied Carbon Coefficients"].iter_rows(values_only=True))
            header_index, headers = self._factor_header(rows, "material category")
            category_col = self._factor_column(headers, ["material category", "category"])
            subtype_col = self._factor_column(headers, ["material sub-type", "sub-type", "subtype"])
            factor_col = self._factor_column(
                headers,
                ["selected a1", "ice ec", "kg co2e/kg", "original factor", "factor"],
            )
            unit_col = self._factor_column(headers, ["unit"])
            source_col = self._factor_column(
                headers,
                ["original source / database", "data source", "source"],
            )
            year_col = self._factor_column(headers, ["factor year", "year"])
            boundary_col = self._factor_column(headers, ["life-cycle boundary", "boundary"])
            url_col = self._factor_column(headers, ["source url", "url"])
            note_col = self._factor_column(
                headers,
                ["verification / replacement note", "notes", "note"],
            )
            geography_col = self._factor_column(headers, ["geography", "location", "region", "country"])
            proxy_col = self._factor_column(headers, ["proxy status", "proxy"])
            for excel_row, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
                category = self._factor_cell(row, category_col)
                subtype = self._factor_cell(row, subtype_col)
                factor = safe_float(self._factor_value(row, factor_col), default=-1.0)
                if category and factor >= 0:
                    self.material_factors.append(
                        {
                            "keyword": f"{category} {subtype}".strip(),
                            "category": category,
                            "subtype": subtype,
                            "factor": factor,
                            "unit": self._factor_cell(row, unit_col),
                            "source": self._factor_cell(row, source_col) or "material factor table",
                            "year": self._factor_cell(row, year_col),
                            "boundary": self._factor_cell(row, boundary_col),
                            "sourceUrl": self._factor_cell(row, url_col),
                            "sourceNote": self._factor_cell(row, note_col),
                            "sourceRowId": self._source_row_id(
                                "Embodied Carbon Coefficients", excel_row
                            ),
                            "geography": self._factor_cell(row, geography_col),
                            "proxyStatus": self._factor_cell(row, proxy_col),
                        }
                    )

        if "Energy Emission Factors" in wb.sheetnames:
            rows = list(wb["Energy Emission Factors"].iter_rows(values_only=True))
            header_index, headers = self._factor_header(rows, "category")
            category_col = self._factor_column(headers, ["category"])
            energy_col = self._factor_column(
                headers,
                ["sub-type / year", "energy type", "sub-type", "type"],
            )
            factor_col = self._factor_column(
                headers,
                ["selected factor", "emission factor", "original factor", "factor"],
            )
            unit_col = self._factor_column(headers, ["unit"])
            year_col = self._factor_column(headers, ["factor year", "year"])
            source_col = self._factor_column(
                headers,
                ["original source / database", "data source", "source"],
            )
            boundary_col = self._factor_column(headers, ["life-cycle boundary", "boundary"])
            url_col = self._factor_column(headers, ["source url", "url"])
            note_col = self._factor_column(
                headers,
                ["selection / boundary note", "notes", "note"],
            )
            geography_col = self._factor_column(headers, ["geography", "location", "region", "country"])
            proxy_col = self._factor_column(headers, ["proxy status", "proxy"])
            for excel_row, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
                category = self._factor_cell(row, category_col)
                energy_type = self._factor_cell(row, energy_col)
                factor = safe_float(self._factor_value(row, factor_col), default=-1.0)
                if energy_type and factor >= 0:
                    self.energy_factors.append(
                        {
                            "keyword": f"{category} {energy_type}".strip(),
                            "category": category,
                            "energyType": energy_type,
                            "factor": factor,
                            "unit": self._factor_cell(row, unit_col),
                            "year": self._factor_cell(row, year_col),
                            "source": self._factor_cell(row, source_col) or "energy factor table",
                            "boundary": self._factor_cell(row, boundary_col),
                            "sourceUrl": self._factor_cell(row, url_col),
                            "sourceNote": self._factor_cell(row, note_col),
                            "sourceRowId": self._source_row_id("Energy Emission Factors", excel_row),
                            "geography": self._factor_cell(row, geography_col),
                            "proxyStatus": self._factor_cell(row, proxy_col),
                        }
                    )

        self.material_factors.sort(key=lambda x: len(normalize_text(x["keyword"])), reverse=True)
        self.energy_factors.sort(key=lambda x: len(normalize_text(x["keyword"])), reverse=True)

    def _source_row_id(self, sheet_name: str, excel_row: int) -> str:
        """Stable identity of the actual source row, independent of candidate order."""
        return f"sha256:{self.workbook_sha256}:{sheet_name}:{excel_row}"

    @staticmethod
    def _factor_header(rows: List[Tuple[Any, ...]], primary_label: str) -> Tuple[int, List[str]]:
        primary = normalize_text(primary_label)
        for index, row in enumerate(rows[:20]):
            headers = [str(value or "").strip() for value in row]
            normalized = [normalize_text(value) for value in headers]
            if any(primary in value for value in normalized) and any("factor" in value for value in normalized):
                return index, headers
        return 0, [str(value or "").strip() for value in rows[0]] if rows else []

    @staticmethod
    def _factor_column(headers: Sequence[str], candidates: Sequence[str]) -> Optional[int]:
        normalized_headers = [normalize_text(value) for value in headers]
        for candidate in candidates:
            normalized_candidate = normalize_text(candidate)
            for index, header in enumerate(normalized_headers):
                if normalized_candidate and normalized_candidate in header:
                    return index
        return None

    @staticmethod
    def _factor_value(row: Sequence[Any], column: Optional[int]) -> Any:
        if column is None or column >= len(row):
            return None
        return row[column]

    @classmethod
    def _factor_cell(cls, row: Sequence[Any], column: Optional[int]) -> str:
        value = cls._factor_value(row, column)
        if value is None:
            return ""
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value).strip()

    @staticmethod
    def _factor_scope_status(factor_unit: str, matched_status: str) -> str:
        unit = normalize_text(factor_unit)
        if "kgco2/" in unit and "kgco2e/" not in unit:
            return "factor_scope_incompatible"
        return matched_status

    @staticmethod
    def _proxy_status(row: Dict[str, Any]) -> str:
        explicit = str(row.get("proxyStatus", "") or "").strip()
        if explicit:
            return explicit
        note = str(row.get("sourceNote", "") or "").casefold()
        return "proxy" if "proxy" in note else "not_proxy"

    @staticmethod
    def _factor_row_id(row: Dict[str, Any], factor_kind: str) -> str:
        source_row_id = str(row.get("sourceRowId", "") or "").strip()
        if source_row_id:
            return source_row_id
        identity = row.get("energyType") or row.get("keyword") or row.get("category") or "unknown"
        return f"factor:{factor_kind}:{slugify(str(identity))}_{safe_float(row.get('factor'))}"

    def _strict_record(self, row: Dict[str, Any], factor_kind: str) -> FactorRecord:
        original_factor_unit = str(row.get("unit", "") or "").strip()
        denominator = parse_factor_denominator(original_factor_unit)
        normalized_factor_value = safe_float(row.get("factor")) / denominator.to_canonical
        source_row_id = self._factor_row_id(row, factor_kind)
        return FactorRecord(
            factor_id=source_row_id,
            keyword=str(row.get("energyType") or row.get("keyword") or row.get("category") or ""),
            factor_value=safe_float(row.get("factor")),
            factor_unit=original_factor_unit,
            source=str(row.get("source", "") or ""),
            match_status="accepted_strict_v2",
            confidence=1.0,
            source_row_id=source_row_id,
            original_factor_unit=original_factor_unit,
            denominator_unit=denominator.canonical_unit,
            normalized_factor_value=normalized_factor_value,
            normalized_denominator=denominator.canonical_unit,
            geography=str(row.get("geography", "") or ""),
            year=str(row.get("year", "") or ""),
            system_boundary=str(row.get("boundary", "") or ""),
            source_reference=str(row.get("sourceUrl") or row.get("source") or ""),
            proxy_status=self._proxy_status(row),
            validation_status="accepted",
        )

    @staticmethod
    def _identity_matches(candidate: str, requested: str) -> bool:
        candidate_norm = normalize_text(candidate)
        requested_norm = normalize_text(requested)
        if not candidate_norm or not requested_norm:
            return False
        compact_candidate = re.sub(r"[\W_]+", "", candidate_norm)
        compact_requested = re.sub(r"[\W_]+", "", requested_norm)
        return candidate_norm == requested_norm or compact_candidate == compact_requested

    def _strict_resolve(
        self,
        candidates: Iterable[Dict[str, Any]],
        factor_kind: str,
        quantity_unit: str,
        quantity_scope: str,
        requested_scope: str,
    ) -> Optional[FactorRecord]:
        rows = list(candidates)
        if not rows:
            self.last_strict_resolution_reason = "unresolved_factor"
            return None
        accepted: List[FactorRecord] = []
        reasons: List[str] = []
        for row in rows:
            validation = validate_quantity_factor(
                quantity_unit,
                str(row.get("unit", "") or ""),
                quantity_scope,
                str(row.get("boundary", "") or ""),
                requested_scope,
            )
            if not validation.accepted:
                reasons.append(validation.reason_code)
                continue
            try:
                accepted.append(self._strict_record(row, factor_kind))
            except UnitError:
                reasons.append("factor_denominator_invalid")
        if len(accepted) == 1:
            self.last_strict_resolution_reason = "accepted"
            return accepted[0]
        if len(accepted) > 1:
            self.last_strict_resolution_reason = "ambiguous_factor"
            return None
        self.last_strict_resolution_reason = reasons[0] if reasons else "unresolved_factor"
        return None

    def resolve_energy_factor_strict_v2(
        self,
        carrier: str,
        quantity_unit: str,
        *,
        quantity_scope: str,
        requested_scope: str,
    ) -> Optional[FactorRecord]:
        """Return one v2-compatible energy factor or ``None``; never a zero placeholder."""
        candidates = [
            row
            for row in self.energy_factors
            if self._identity_matches(str(row.get("energyType", "")), carrier)
            or self._identity_matches(str(row.get("keyword", "")), carrier)
        ]
        return self._strict_resolve(
            candidates, "energy", quantity_unit, quantity_scope, requested_scope
        )

    def resolve_material_factor_strict_v2(
        self,
        material_name: str,
        quantity_unit: str,
        *,
        quantity_scope: str,
        requested_scope: str,
    ) -> Optional[FactorRecord]:
        """Return one v2-compatible material factor or ``None``; never a zero placeholder."""
        candidates = [
            row
            for row in self.material_factors
            if any(
                self._identity_matches(str(row.get(key, "")), material_name)
                for key in ("category", "subtype", "keyword")
            )
        ]
        return self._strict_resolve(
            candidates, "material", quantity_unit, quantity_scope, requested_scope
        )

    # Explicit aliases make v2-only consumers readable while legacy methods remain staged.
    resolve_energy_factor_v2 = resolve_energy_factor_strict_v2
    resolve_material_factor_v2 = resolve_material_factor_strict_v2

    def resolve_material_factor(self, material_name: str) -> FactorRecord:
        norm = normalize_text(material_name)
        family, _density = keyword_family(material_name)

        for row in self.material_factors:
            tokens = [row.get("category", ""), row.get("subtype", ""), row.get("keyword", "")]
            for token in tokens:
                token_norm = normalize_text(token)
                if token_norm and token_norm in norm:
                    factor_unit = row.get("unit", "kgCO2e/kg")
                    return FactorRecord(
                        factor_id=f"factor:material:{slugify(row['keyword'])}_{row['factor']}",
                        keyword=row["keyword"],
                        factor_value=safe_float(row["factor"]),
                        factor_unit=factor_unit,
                        source=row.get("source", "factor table"),
                        match_status=self._factor_scope_status(factor_unit, "exact_or_keyword"),
                        confidence=0.9,
                    )

        if family != "unknown":
            for row in self.material_factors:
                if normalize_text(family) in normalize_text(row.get("category", "")):
                    factor_unit = row.get("unit", "kgCO2e/kg")
                    return FactorRecord(
                        factor_id=f"factor:material:{slugify(row['keyword'])}_{row['factor']}",
                        keyword=row["keyword"],
                        factor_value=safe_float(row["factor"]),
                        factor_unit=factor_unit,
                        source=row.get("source", "factor table"),
                        match_status=self._factor_scope_status(factor_unit, "family_fallback"),
                        confidence=0.65,
                    )

        return FactorRecord(
            factor_id="factor:material:unresolved",
            keyword="unresolved",
            factor_value=0.0,
            factor_unit="kgCO2e/kg",
            source="unresolved",
            match_status="unresolved",
            confidence=0.0,
        )

    def resolve_energy_factor(self, carrier: str, quantity_unit: str = "") -> FactorRecord:
        norm = normalize_text(carrier)
        compact_norm = re.sub(r"[\W_]+", "", norm)
        unit_norm = normalize_text(quantity_unit)
        candidates: List[Dict[str, Any]] = []
        for row in self.energy_factors:
            keyword_norm = normalize_text(row.get("keyword", ""))
            energy_norm = normalize_text(row.get("energyType", ""))
            compact_keyword = re.sub(r"[\W_]+", "", keyword_norm)
            compact_energy = re.sub(r"[\W_]+", "", energy_norm)
            if (
                (energy_norm and energy_norm in norm)
                or (keyword_norm and norm in keyword_norm)
                or (compact_energy and compact_energy in compact_norm)
                or (compact_keyword and compact_norm in compact_keyword)
            ):
                candidates.append(row)

        if unit_norm and candidates:
            compatible = []
            for row in candidates:
                factor_unit_norm = normalize_text(row.get("unit", ""))
                if unit_norm in factor_unit_norm:
                    compatible.append(row)
            if compatible:
                candidates = compatible
            else:
                return FactorRecord(
                    factor_id="factor:energy:unit_incompatible",
                    keyword=carrier,
                    factor_value=0.0,
                    factor_unit="unresolved",
                    source="unresolved",
                    match_status="unit_incompatible",
                    confidence=0.0,
                )

        if candidates:
            def electricity_priority(row: Mapping[str, Any]) -> tuple[int, str]:
                text = normalize_text(
                    " ".join(
                        str(row.get(key, "") or "")
                        for key in ("category", "energyType", "keyword", "sourceNote", "geography")
                    )
                )
                score = 0
                if "guangdong lifecycle co2e proxy" in text:
                    score += 200
                if "guangdong factory proxy" in text:
                    score += 100
                if "lifecycle" in text or "carbon footprint" in text:
                    score += 20
                return score, str(row.get("sourceRowId", ""))

            row = max(candidates, key=electricity_priority) if norm == "electricity" else candidates[0]
            factor_unit = row.get("unit", "kgCO2e/kWh")
            return FactorRecord(
                factor_id=f"factor:energy:{slugify(row['energyType'])}_{row['factor']}",
                keyword=row.get("energyType", ""),
                factor_value=safe_float(row.get("factor")),
                factor_unit=factor_unit,
                source=row.get("source", "energy factor table"),
                match_status=self._factor_scope_status(factor_unit, "exact_or_keyword"),
                confidence=0.9,
            )

        return FactorRecord(
            factor_id="factor:energy:unresolved",
            keyword="unresolved",
            factor_value=0.0,
            factor_unit="kgCO2e/kWh",
            source="unresolved",
            match_status="unresolved",
            confidence=0.0,
        )


class IFCDesignExtractor:
    def __init__(self, ifc_path: Path, include_openings: bool = False) -> None:
        if ifcopenshell is None:
            raise RuntimeError("ifcopenshell is required to parse IFC")
        self.ifc_path = ifc_path
        self.ifc = ifcopenshell.open(str(ifc_path))
        self.include_openings = include_openings
        self.material_assoc: Dict[str, List[MaterialRecord]] = defaultdict(list)
        self.type_material_assoc: Dict[str, List[MaterialRecord]] = defaultdict(list)
        self.quantities: Dict[str, Dict[str, float]] = defaultdict(dict)
        self.storeys: Dict[str, str] = {}
        self.group_names: Dict[str, List[str]] = defaultdict(list)
        self.type_map: Dict[str, Dict[str, str]] = {}
        self._build_caches()

    @staticmethod
    def _entity_key(entity: Any) -> str:
        gid = str(getattr(entity, "GlobalId", "") or "")
        if gid:
            return gid
        ent_id = getattr(entity, "id", None)
        if callable(ent_id):
            try:
                return f"{entity.is_a()}::{ent_id()}"
            except Exception:
                pass
        return ""

    @staticmethod
    def _clone_material(row: MaterialRecord, source_suffix: str) -> MaterialRecord:
        return MaterialRecord(
            name=row.name,
            source=f"{row.source}:{source_suffix}",
            thickness=row.thickness,
            status=row.status,
        )

    @staticmethod
    def _unique_materials(rows: List[MaterialRecord]) -> List[MaterialRecord]:
        unique: List[MaterialRecord] = []
        seen: Set[Tuple[str, str, Optional[float]]] = set()
        for row in rows:
            key = (row.name, row.source, row.thickness)
            if key in seen:
                continue
            seen.add(key)
            unique.append(row)
        return unique

    def _build_caches(self) -> None:
        for rel in self.ifc.by_type("IfcRelAssociatesMaterial"):
            materials = self._extract_materials(getattr(rel, "RelatingMaterial", None))
            for obj in getattr(rel, "RelatedObjects", []) or []:
                key = self._entity_key(obj)
                if not key:
                    continue
                if obj.is_a("IfcTypeObject"):
                    self.type_material_assoc[key].extend(
                        [self._clone_material(m, "type_object") for m in materials]
                    )
                else:
                    self.material_assoc[key].extend(materials)

        for rel in self.ifc.by_type("IfcRelDefinesByProperties"):
            pset = getattr(rel, "RelatingPropertyDefinition", None)
            if pset is None or not pset.is_a("IfcElementQuantity"):
                continue
            for obj in getattr(rel, "RelatedObjects", []) or []:
                gid = getattr(obj, "GlobalId", None)
                if not gid:
                    continue
                for qty in getattr(pset, "Quantities", []) or []:
                    value = self._quantity_value(qty)
                    if value is not None:
                        self.quantities[gid][getattr(qty, "Name", qty.is_a())] = value

        for rel in self.ifc.by_type("IfcRelContainedInSpatialStructure"):
            storey_name = getattr(getattr(rel, "RelatingStructure", None), "Name", "") or ""
            for elem in getattr(rel, "RelatedElements", []) or []:
                gid = getattr(elem, "GlobalId", None)
                if gid:
                    self.storeys[gid] = storey_name

        for rel in self.ifc.by_type("IfcRelAssignsToGroup"):
            group_name = getattr(getattr(rel, "RelatingGroup", None), "Name", "") or ""
            for obj in getattr(rel, "RelatedObjects", []) or []:
                gid = getattr(obj, "GlobalId", None)
                if gid and group_name:
                    self.group_names[gid].append(group_name)

        for rel in self.ifc.by_type("IfcRelDefinesByType"):
            type_obj = getattr(rel, "RelatingType", None)
            if type_obj is None:
                continue
            type_key = self._entity_key(type_obj)
            if type_key:
                for assoc in getattr(type_obj, "HasAssociations", []) or []:
                    if assoc is None or not assoc.is_a("IfcRelAssociatesMaterial"):
                        continue
                    materials = self._extract_materials(getattr(assoc, "RelatingMaterial", None))
                    self.type_material_assoc[type_key].extend(
                        [self._clone_material(m, "type_object_association") for m in materials]
                    )

            for obj in getattr(rel, "RelatedObjects", []) or []:
                gid = getattr(obj, "GlobalId", None)
                if gid:
                    self.type_map[gid] = {
                        "typeObjectName": str(getattr(type_obj, "Name", "") or ""),
                        "typeObjectIfcType": str(type_obj.is_a()),
                        "typeObjectGlobalId": str(getattr(type_obj, "GlobalId", "") or ""),
                        "typeObjectKey": type_key,
                    }

    def _extract_materials(self, material_select: Any) -> List[MaterialRecord]:
        rows: List[MaterialRecord] = []

        def add(name: Any, source: str, thickness: Any = None) -> None:
            text = str(name or "").strip()
            if not text:
                return
            rows.append(
                MaterialRecord(
                    name=text,
                    source=source,
                    thickness=safe_float(thickness) if thickness not in (None, "") else None,
                    status="ambiguous" if is_placeholder_material(text) else "resolved",
                )
            )

        mat = material_select
        if mat is None:
            return rows
        if mat.is_a("IfcMaterial"):
            add(getattr(mat, "Name", ""), "IfcMaterial")
        elif mat.is_a("IfcMaterialLayerSetUsage"):
            for layer in getattr(getattr(mat, "ForLayerSet", None), "MaterialLayers", []) or []:
                add(getattr(getattr(layer, "Material", None), "Name", ""), "IfcMaterialLayerSetUsage", getattr(layer, "LayerThickness", None))
        elif mat.is_a("IfcMaterialLayerSet"):
            for layer in getattr(mat, "MaterialLayers", []) or []:
                add(getattr(getattr(layer, "Material", None), "Name", ""), "IfcMaterialLayerSet", getattr(layer, "LayerThickness", None))
        elif mat.is_a("IfcMaterialConstituentSet"):
            for item in getattr(mat, "MaterialConstituents", []) or []:
                add(getattr(getattr(item, "Material", None), "Name", ""), "IfcMaterialConstituentSet")
        elif mat.is_a("IfcMaterialProfileSetUsage"):
            for item in getattr(getattr(mat, "ForProfileSet", None), "MaterialProfiles", []) or []:
                add(getattr(getattr(item, "Material", None), "Name", ""), "IfcMaterialProfileSetUsage")
        elif mat.is_a("IfcMaterialProfileSet"):
            for item in getattr(mat, "MaterialProfiles", []) or []:
                add(getattr(getattr(item, "Material", None), "Name", ""), "IfcMaterialProfileSet")
        elif mat.is_a("IfcMaterialList"):
            for item in getattr(mat, "Materials", []) or []:
                add(getattr(item, "Name", ""), "IfcMaterialList")

        deduped: List[MaterialRecord] = []
        seen: Set[Tuple[str, str, Optional[float]]] = set()
        for row in rows:
            key = (row.name, row.source, row.thickness)
            if key not in seen:
                seen.add(key)
                deduped.append(row)
        return deduped

    @staticmethod
    def _quantity_value(qty: Any) -> Optional[float]:
        for cls, attr in [
            ("IfcQuantityLength", "LengthValue"),
            ("IfcQuantityArea", "AreaValue"),
            ("IfcQuantityVolume", "VolumeValue"),
            ("IfcQuantityWeight", "WeightValue"),
            ("IfcQuantityCount", "CountValue"),
        ]:
            if qty.is_a(cls):
                return round(safe_float(getattr(qty, attr, None)), 6)
        return None

    def extract_components(self, element_types: Optional[List[str]] = None) -> List[ComponentRecord]:
        rows: List[ComponentRecord] = []
        seen: Set[str] = set()
        if element_types is None:
            elements = self.ifc.by_type("IfcElement")
        else:
            elements = []
            for ifc_type in element_types:
                if ifc_type == "IfcOpeningElement" and not self.include_openings:
                    continue
                try:
                    elements.extend(self.ifc.by_type(ifc_type))
                except Exception:
                    continue

        for entity in elements:
            if entity.is_a("IfcOpeningElement") and not self.include_openings:
                continue
            gid = getattr(entity, "GlobalId", None)
            if not gid or gid in seen:
                continue
            seen.add(gid)
            rows.append(
                ComponentRecord(
                    global_id=gid,
                    ifc_type=entity.is_a(),
                    name=getattr(entity, "Name", "") or "",
                    object_type=getattr(entity, "ObjectType", "") or "",
                    predefined_type=str(getattr(entity, "PredefinedType", "") or ""),
                    tag=getattr(entity, "Tag", "") or "",
                    storey=self.storeys.get(gid, ""),
                    group_names=sorted(set(self.group_names.get(gid, []))),
                    materials=self._unique_materials(
                        self.material_assoc.get(gid, [])
                        + self.type_material_assoc.get(self.type_map.get(gid, {}).get("typeObjectKey", ""), [])
                    ),
                    quantities=self.quantities.get(gid, {}),
                )
            )
        return rows


class MultiGranularCarbonKGBuilder:
    def __init__(
        self,
        ifc_path: Path,
        factor_library: FactorLibrary,
        module_name: str,
        include_openings: bool = False,
        energy_csv: Optional[Path] = None,
        factory_xlsx: Optional[Path] = None,
        factory_target_map: Optional[Path] = None,
        extraction_result: Optional[IFCExtractionResult] = None,
        accepted_materials: Iterable[AcceptedConsumption] = (),
        accepted_energies: Iterable[
            tuple[AcceptedConsumption, EnergyPopulationPlan]
        ] = (),
        module_identity: Optional[str] = None,
    ) -> None:
        self.ifc_path = ifc_path
        self.factor_library = factor_library
        self.module_name = module_name
        self.include_openings = include_openings
        # The legacy CLI still supplies these arguments until the Task 8 cutover.
        # Canonical v2 construction deliberately does not retain or interpret them.
        _ = (energy_csv, factory_xlsx, factory_target_map)
        self.extraction_result = extraction_result
        self.accepted_materials = tuple(accepted_materials)
        self.accepted_energies = tuple(accepted_energies)
        self.module_identity = str(module_identity).strip() if module_identity is not None else ""
        self.graph = CanonicalLPGGraph()
        self.stats = Counter()
        self.issues: List[Dict[str, Any]] = []
        self.component_node_by_gid: Dict[str, str] = {}
        self.module_id: Optional[str] = None

    def build(self) -> CanonicalLPGGraph:
        self.build_product_backbone()
        for accepted in self.accepted_materials:
            self.populate_accepted_material(accepted)
        for accepted, plan in self.accepted_energies:
            self.populate_accepted_energy(accepted, plan)
        return self.graph

    @staticmethod
    def _component_props(component: CanonicalComponentRecord) -> Dict[str, Any]:
        return {
            "ifcSha256": component.ifc_hash,
            "globalId": component.global_id,
            "stepId": component.step_id,
            "ifcClass": component.ifc_class,
            "name": component.name,
            "description": component.description,
            "objectType": component.object_type,
            "predefinedType": component.predefined_type,
            "tag": component.tag,
        }

    @staticmethod
    def _component_type_props(component_type: ComponentTypeRecord) -> Dict[str, Any]:
        return {
            "ifcSha256": component_type.ifc_hash,
            "globalId": component_type.global_id,
            "stepId": component_type.step_id,
            "ifcClass": component_type.ifc_class,
            "name": component_type.name,
            "description": component_type.description,
            "tag": component_type.tag,
        }

    @staticmethod
    def _material_edge_props(association: MaterialAssociationRecord) -> Dict[str, Any]:
        return {
            "associationRecordId": association.id,
            "associationStepId": association.association_step_id,
            "sourceKind": association.source_kind,
            "containerKind": association.container_kind,
            "containerStepId": association.container_step_id,
            "itemStepId": association.item_step_id,
            "layerOrConstituentIndex": association.layer_or_constituent_index,
            "thickness": association.thickness,
            "thicknessUnit": association.thickness_unit,
            "constituentFraction": association.constituent_fraction,
        }

    @staticmethod
    def _design_quantity_props(quantity: DesignQuantityRecord) -> Dict[str, Any]:
        return {
            "ifcSha256": quantity.ifc_hash,
            "componentId": quantity.component_id,
            "quantityStepId": quantity.quantity_step_id,
            "qtoSetStepId": quantity.qto_set_step_id,
            "qtoSetName": quantity.qto_set_name,
            "quantityName": quantity.quantity_name,
            "quantitySubtype": quantity.quantity_subtype,
            "rawValue": quantity.source_value,
            "rawUnit": quantity.source_unit,
            "sourceValue": quantity.source_value,
            "sourceUnit": quantity.source_unit,
            "normalizedValue": quantity.normalized_value,
            "normalizedUnit": quantity.normalized_unit,
        }

    @staticmethod
    def _commit_graph(target: CanonicalLPGGraph, staged: CanonicalLPGGraph) -> None:
        target.nodes.clear()
        target.nodes.update(staged.nodes)
        target.edges[:] = staged.edges
        target._edges_by_id.clear()
        target._edges_by_id.update(staged._edges_by_id)

    def build_product_backbone(
        self, extraction: Optional[IFCExtractionResult] = None
    ) -> CanonicalLPGGraph:
        source = extraction or self.extraction_result
        if source is None:
            source = CanonicalIFCExtractor(
                self.ifc_path, include_openings=self.include_openings
            ).extract()
        if not isinstance(source, IFCExtractionResult):
            raise TypeError("extraction must be an IFCExtractionResult")

        target_graph = self.graph
        staged_graph = deepcopy(target_graph)
        previous_extraction = self.extraction_result
        previous_components = deepcopy(self.component_node_by_gid)
        previous_stats = deepcopy(self.stats)
        previous_module_id = self.module_id
        self.graph = staged_graph
        try:
            self._populate_product_backbone(source)
        except Exception:
            self.graph = target_graph
            self.extraction_result = previous_extraction
            self.component_node_by_gid.clear()
            self.component_node_by_gid.update(previous_components)
            self.stats.clear()
            self.stats.update(previous_stats)
            self.module_id = previous_module_id
            raise
        self.graph = target_graph
        self._commit_graph(target_graph, staged_graph)
        return target_graph

    def _populate_product_backbone(self, source: IFCExtractionResult) -> None:
        self.extraction_result = source

        if self.module_identity:
            self.module_id = stable_id(
                "ModularUnit", source.ifc_sha256, self.module_identity
            )
            self.graph.add_node(
                self.module_id,
                ["ModularUnit"],
                {
                    "ifcSha256": source.ifc_sha256,
                    "sourceIdentity": self.module_identity,
                    "name": self.module_name,
                },
            )

        for component in sorted(source.components, key=lambda row: row.id):
            self.graph.add_node(
                component.id,
                ["BuildingComponent", component.ifc_class],
                self._component_props(component),
            )
            self.component_node_by_gid[component.global_id] = component.id
            if self.module_id is not None:
                self.graph.add_edge(
                    self.module_id,
                    "containsComponent",
                    component.id,
                    {"sourceIdentity": self.module_identity},
                    occurrence_id=stable_id(
                        "ModuleComponentAssociation",
                        source.ifc_sha256,
                        self.module_identity,
                        component.id,
                    ),
                )
            if component.component_type is not None:
                component_type = component.component_type
                self.graph.add_node(
                    component_type.id,
                    ["ComponentType", component_type.ifc_class],
                    self._component_type_props(component_type),
                )
                self.graph.add_edge(
                    component.id,
                    "hasComponentType",
                    component_type.id,
                    occurrence_id=stable_id(
                        "ComponentTypeAssignment", component.id, component_type.id
                    ),
                )

        for association in sorted(
            deduplicate_material_associations(source.material_associations),
            key=lambda row: row.id,
        ):
            if association.component_id not in self.graph.nodes:
                raise ValueError(
                    f"material association {association.id!r} references missing component "
                    f"{association.component_id!r}"
                )
            self.graph.add_node(
                association.material_id,
                ["IfcMaterial"],
                {
                    "ifcSha256": association.ifc_hash,
                    "materialStepId": association.material_step_id,
                    "name": association.material_name,
                },
            )
            self.graph.add_edge(
                association.component_id,
                "hasMaterial",
                association.material_id,
                self._material_edge_props(association),
                occurrence_id=association.id,
            )

        for quantity in sorted(source.design_quantities, key=lambda row: row.id):
            if quantity.component_id not in self.graph.nodes:
                raise ValueError(
                    f"design quantity {quantity.id!r} references missing component "
                    f"{quantity.component_id!r}"
                )
            self.graph.add_node(
                quantity.id,
                ["DesignQuantity", quantity.quantity_subtype],
                self._design_quantity_props(quantity),
            )
            self.graph.add_edge(
                quantity.component_id,
                "hasDesignQuantity",
                quantity.id,
                occurrence_id=quantity.id,
            )

        self.stats["componentsTotal"] = len(source.components)
        self.stats["materialsTotal"] = self.graph.count_label("IfcMaterial")
        self.stats["materialAssociationsTotal"] = self.graph.count_relation("hasMaterial")
        self.stats["designQuantitiesTotal"] = self.graph.count_label("DesignQuantity")

    def populate_accepted_material(self, accepted: AcceptedConsumption) -> None:
        if not isinstance(accepted, AcceptedConsumption):
            raise TypeError("accepted must be an AcceptedConsumption")
        if accepted.kind != "material":
            raise ValueError("populate_accepted_material requires a material fact")
        factor = accepted.factor
        if (
            not accepted.factor_source_id
            or accepted.factor_source_id != factor.source_row_id
        ):
            raise ValueError(
                "accepted.factor_source_id must equal accepted.factor.source_row_id"
            )
        factor_node_id = stable_id("EmissionFactor", accepted.factor_source_id)
        if not accepted.product_target_id or accepted.product_target_id not in self.graph.nodes:
            raise ValueError(
                f"accepted material {accepted.record_id!r} references missing BuildingComponent "
                f"{accepted.product_target_id!r}"
            )
        if "BuildingComponent" not in self.graph.nodes[accepted.product_target_id]["labels"]:
            raise ValueError(
                f"accepted material {accepted.record_id!r} target is not a BuildingComponent"
            )
        if not accepted.material_id or accepted.material_id not in self.graph.nodes:
            raise ValueError(
                f"accepted material {accepted.record_id!r} references missing IfcMaterial "
                f"{accepted.material_id!r}"
            )
        if "IfcMaterial" not in self.graph.nodes[accepted.material_id]["labels"]:
            raise ValueError(
                f"accepted material {accepted.record_id!r} material target is not an IfcMaterial"
            )
        design_quantity_ids = tuple(
            dict.fromkeys(
                operand.design_quantity_id
                for operand in accepted.operands
                if operand.design_quantity_id
            )
        )
        missing_quantities = [
            quantity_id
            for quantity_id in design_quantity_ids
            if quantity_id not in self.graph.nodes
            or "DesignQuantity" not in self.graph.nodes[quantity_id]["labels"]
        ]
        if missing_quantities:
            raise ValueError(
                f"accepted material {accepted.record_id!r} references missing DesignQuantity "
                f"{missing_quantities!r}"
            )
        material_edges = [
            edge
            for edge in self.graph.edges_of_type("hasMaterial")
            if edge["src"] == accepted.product_target_id
            and edge["tgt"] == accepted.material_id
        ]
        if not material_edges:
            raise ValueError(
                f"accepted material {accepted.record_id!r} has no hasMaterial ownership "
                "between its BuildingComponent and IfcMaterial"
            )
        if not any(
            edge["occurrenceId"] == accepted.source_record_id
            for edge in material_edges
        ):
            raise ValueError(
                f"accepted material {accepted.record_id!r} source_record_id does not identify "
                "the exact hasMaterial association occurrence"
            )
        owned_design_quantities = {
            edge["tgt"]
            for edge in self.graph.edges_of_type("hasDesignQuantity")
            if edge["src"] == accepted.product_target_id
        }
        unowned_quantities = [
            quantity_id
            for quantity_id in design_quantity_ids
            if quantity_id not in owned_design_quantities
        ]
        if unowned_quantities:
            raise ValueError(
                f"accepted material {accepted.record_id!r} DesignQuantity provenance lacks "
                f"hasDesignQuantity ownership: {unowned_quantities!r}"
            )

        common_props = {
            "recordId": accepted.record_id,
            "sourceIdentity": accepted.source_identity,
            "sourceRecordId": accepted.source_record_id,
            "evidenceSourceId": accepted.evidence_source_id,
            "formulaCode": accepted.formula_code,
            "recordedScope": accepted.recorded_scope,
            "requestedScope": accepted.requested_scope,
            "systemBoundary": accepted.system_boundary,
            "isValidZero": accepted.is_valid_zero,
            "measured": accepted.measured,
            "dataProvenance": accepted.data_provenance,
        }
        consumption_props = {
            **common_props,
            "productTargetId": accepted.product_target_id,
            "materialId": accepted.material_id,
            "quantityId": accepted.quantity_id,
            "factorId": factor_node_id,
            "factorSourceId": accepted.factor_source_id,
            "factorAliasId": accepted.factor_id,
            "emissionId": accepted.emission_id,
        }
        quantity_props = {
            **common_props,
            "consumptionId": accepted.consumption_id,
            "orderedRawOperands": [
                json.dumps(
                    asdict(operand),
                    ensure_ascii=False,
                    sort_keys=True,
                    separators=(",", ":"),
                )
                for operand in accepted.operands
            ],
            "conversionSteps": [
                json.dumps(
                    asdict(step),
                    ensure_ascii=False,
                    sort_keys=True,
                    separators=(",", ":"),
                )
                for step in accepted.conversion_steps
            ],
            "quantityValue": accepted.quantity_value,
            "quantityUnit": accepted.quantity_unit,
            "quantityConversionFactor": accepted.quantity_conversion_factor,
        }
        factor_props = {
            "factorId": factor_node_id,
            "factorSourceId": accepted.factor_source_id,
            "keyword": factor.keyword,
            "factorValue": accepted.factor_value,
            "factorUnit": factor.factor_unit,
            "factorDenominator": accepted.factor_denominator,
            "originalFactorValue": factor.factor_value,
            "originalFactorUnit": factor.original_factor_unit,
            "source": factor.source,
            "sourceRowId": factor.source_row_id,
            "matchStatus": factor.match_status,
            "confidence": factor.confidence,
            "denominatorUnit": factor.denominator_unit,
            "normalizedFactorValue": factor.normalized_factor_value,
            "normalizedDenominator": factor.normalized_denominator,
            "geography": factor.geography,
            "year": factor.year,
            "systemBoundary": factor.system_boundary,
            "sourceReference": factor.source_reference,
            "proxyStatus": factor.proxy_status,
            "validationStatus": factor.validation_status,
        }
        emission_props = {
            **common_props,
            "consumptionId": accepted.consumption_id,
            "quantityId": accepted.quantity_id,
            "factorId": factor_node_id,
            "factorSourceId": accepted.factor_source_id,
            "factorAliasId": accepted.factor_id,
            "quantityValue": accepted.quantity_value,
            "quantityUnit": accepted.quantity_unit,
            "factorValue": accepted.factor_value,
            "factorDenominator": accepted.factor_denominator,
            "emissionValue": accepted.emission_value,
            "emissionUnit": accepted.emission_unit,
        }
        staged_graph = deepcopy(self.graph)
        staged_graph.add_node(
            accepted.consumption_id,
            ["MaterialConsumption"],
            consumption_props,
        )
        staged_graph.add_node(
            accepted.quantity_id,
            ["ConsumptionQuantity"],
            quantity_props,
        )
        staged_graph.add_node(
            factor_node_id,
            ["EmissionFactor"],
            factor_props,
        )
        staged_graph.add_node(
            accepted.emission_id,
            ["CarbonEmission"],
            emission_props,
        )

        def occurrence(relation: str, target_id: str) -> str:
            return stable_id(
                "MaterialCalculationRelationship",
                accepted.consumption_id,
                relation,
                target_id,
            )

        staged_graph.add_edge(
            accepted.consumption_id,
            "recordedForObject",
            accepted.product_target_id,
            {"recordId": accepted.record_id},
            occurrence_id=occurrence("recordedForObject", accepted.product_target_id),
        )
        staged_graph.add_edge(
            accepted.consumption_id,
            "ofMaterial",
            accepted.material_id,
            {"recordId": accepted.record_id},
            occurrence_id=occurrence("ofMaterial", accepted.material_id),
        )
        staged_graph.add_edge(
            accepted.consumption_id,
            "hasQuantity",
            accepted.quantity_id,
            occurrence_id=occurrence("hasQuantity", accepted.quantity_id),
        )
        staged_graph.add_edge(
            accepted.consumption_id,
            "hasFactor",
            factor_node_id,
            occurrence_id=occurrence("hasFactor", factor_node_id),
        )
        staged_graph.add_edge(
            accepted.emission_id,
            "hasCarbonDriver",
            accepted.consumption_id,
            occurrence_id=stable_id(
                "MaterialCalculationRelationship",
                accepted.emission_id,
                "hasCarbonDriver",
                accepted.consumption_id,
            ),
        )
        for design_quantity_id in design_quantity_ids:
            staged_graph.add_edge(
                accepted.quantity_id,
                "derivedFrom",
                design_quantity_id,
                occurrence_id=stable_id(
                    "MaterialCalculationRelationship",
                    accepted.quantity_id,
                    "derivedFrom",
                    design_quantity_id,
                ),
            )
        self._commit_graph(self.graph, staged_graph)

    def populate_accepted_energy(
        self,
        accepted: AcceptedConsumption,
        plan: EnergyPopulationPlan,
    ) -> None:
        """Populate one validated energy fact and its explicit attribution plan.

        The operation is transactional: validation and all graph writes occur
        against a shadow graph, and the builder graph is committed only after
        the complete canonical fact succeeds.
        """

        if not isinstance(accepted, AcceptedConsumption):
            raise TypeError("accepted must be an AcceptedConsumption")
        if not isinstance(plan, EnergyPopulationPlan):
            raise TypeError("plan must be an EnergyPopulationPlan")
        validate_energy_population_plan(plan)
        if accepted.kind != "energy":
            raise ValueError("populate_accepted_energy requires an energy fact")
        if plan.attribution_mode not in {"direct", "allocated", "process_only"}:
            raise ValueError("unsupported energy attribution_mode")
        factor = accepted.factor
        if (
            not accepted.factor_source_id
            or accepted.factor_source_id != factor.source_row_id
        ):
            raise ValueError(
                "accepted.factor_source_id must equal accepted.factor.source_row_id"
            )
        if not str(accepted.energy_carrier_id or "").strip():
            raise ValueError("accepted energy requires a canonical energy_carrier_id")

        allocation_set_id = str(plan.allocation_set_id or "").strip()
        allocation_basis = str(plan.allocation_basis or "").strip()
        allocations = tuple(plan.allocations)
        process_node_ids = tuple(plan.process_node_ids)
        resource_node_ids = tuple(plan.resource_node_ids)

        if len(set(process_node_ids)) != len(process_node_ids):
            raise ValueError("process_node_ids must be unique")
        if len(set(resource_node_ids)) != len(resource_node_ids):
            raise ValueError("resource_node_ids must be unique")
        for process_node_id in process_node_ids:
            node = self.graph.nodes.get(process_node_id)
            if node is None or not set(node.get("labels", ())) & {
                "ProductionStage",
                "ManufacturingActivity",
            }:
                raise ValueError(
                    f"process context {process_node_id!r} is not an existing canonical process node"
                )
        for resource_node_id in resource_node_ids:
            node = self.graph.nodes.get(resource_node_id)
            if node is None or "ManufacturingResource" not in node.get("labels", ()):
                raise ValueError(
                    f"resource context {resource_node_id!r} is not an existing ManufacturingResource"
                )

        if plan.attribution_mode == "direct":
            if allocations or allocation_set_id or allocation_basis:
                raise ValueError("direct attribution cannot carry allocation metadata")
            target_id = str(accepted.product_target_id or "").strip()
            target = self.graph.nodes.get(target_id)
            if target is None or not (
                "BuildingComponent" in target.get("labels", ())
                or "ModularUnit" in target.get("labels", ())
            ):
                raise ValueError(
                    "direct energy requires an existing ProductObject product_target_id"
                )
        elif plan.attribution_mode == "allocated":
            if accepted.product_target_id is not None:
                raise ValueError(
                    "allocated energy source facts cannot carry product_target_id"
                )
            if not allocation_set_id or not allocation_basis or not allocations:
                raise ValueError(
                    "allocated energy requires allocation_set_id, allocation_basis, and targets"
                )
            target_ids: set[str] = set()
            normalized_total = 0.0
            for allocation in allocations:
                if not isinstance(allocation, EnergyAllocationTarget):
                    raise TypeError("allocations must contain EnergyAllocationTarget records")
                target_id = str(allocation.target_component_id or "").strip()
                if not target_id or target_id in target_ids:
                    raise ValueError("allocation targets must be non-empty and unique")
                target_ids.add(target_id)
                target = self.graph.nodes.get(target_id)
                if target is None or "BuildingComponent" not in target.get("labels", ()):
                    raise ValueError(
                        f"allocation target {target_id!r} is not an existing BuildingComponent"
                    )
                try:
                    raw_weight = float(allocation.raw_weight)
                    normalized_weight = float(allocation.normalized_weight)
                except (TypeError, ValueError) as exc:
                    raise ValueError("allocation weights must be numeric") from exc
                if (
                    not math.isfinite(raw_weight)
                    or raw_weight < 0
                    or not math.isfinite(normalized_weight)
                    or normalized_weight < 0
                    or normalized_weight > 1
                ):
                    raise ValueError(
                        "allocation weights must be finite, non-negative, and allocatedFraction <= 1"
                    )
                if not str(allocation.raw_weight_unit or "").strip():
                    raise ValueError("raw_weight_unit must be non-empty")
                if not str(allocation.evidence_record_id or "").strip():
                    raise ValueError("evidence_record_id must be non-empty")
                normalized_total += normalized_weight
            unattributed_fraction = float(plan.unattributed_fraction)
            if (
                not math.isfinite(unattributed_fraction)
                or unattributed_fraction < 0.0
                or unattributed_fraction > 1.0
                or not math.isclose(
                    normalized_total + unattributed_fraction,
                    1.0,
                    rel_tol=0.0,
                    abs_tol=1e-9,
                )
            ):
                raise ValueError(
                    "supplied allocated and unattributed fractions must sum to 1; implicit normalization is forbidden"
                )
        else:
            if accepted.product_target_id is not None:
                raise ValueError("process-only energy cannot carry product_target_id")
            if allocations or allocation_set_id or allocation_basis:
                raise ValueError("process-only energy cannot carry allocation metadata")
            if not process_node_ids and not resource_node_ids:
                raise ValueError(
                    "process-only energy requires an existing process or resource context"
                )

        planned_product_edge_props: Dict[str, Dict[str, Any]] = {}
        if plan.attribution_mode == "direct":
            assert accepted.product_target_id is not None
            planned_product_edge_props[accepted.product_target_id] = {
                "sourceRecordId": accepted.source_record_id,
                "evidenceSourceId": accepted.evidence_source_id,
            }
        elif plan.attribution_mode == "allocated":
            planned_product_edge_props = {
                allocation.target_component_id: {
                    "allocated": True,
                    "allocationSetId": allocation_set_id,
                    "allocationBasis": allocation_basis,
                    "rawWeight": float(allocation.raw_weight),
                    "rawWeightUnit": allocation.raw_weight_unit,
                    "allocatedFraction": float(allocation.normalized_weight),
                    "evidenceRecordId": allocation.evidence_record_id,
                }
                for allocation in allocations
            }

        existing_product_edges = [
            edge
            for edge in self.graph.edges_of_type("recordedForObject")
            if edge["src"] == accepted.consumption_id
        ]
        if existing_product_edges and {
            edge["tgt"] for edge in existing_product_edges
        } != set(planned_product_edge_props):
            raise ValueError(
                "existing product attribution target set conflicts with the incoming energy plan"
            )
        if plan.attribution_mode == "process_only" and existing_product_edges:
            raise ValueError("process-only energy cannot reuse an existing product path")
        for edge in existing_product_edges:
            target_id = edge["tgt"]
            expected_props = planned_product_edge_props.get(target_id)
            if expected_props is None:
                raise ValueError(
                    "existing product attribution conflicts with the incoming energy plan"
                )
            if plan.attribution_mode == "allocated":
                edge_props = edge.get("props", {})
                if (
                    edge_props.get("allocated") is not True
                    or edge_props.get("allocationSetId") != allocation_set_id
                    or edge_props != expected_props
                ):
                    raise ValueError(
                        "conflicting allocation evidence for existing "
                        f"({accepted.emission_id!r}, {allocation_set_id!r}, {target_id!r})"
                    )
            elif edge.get("props", {}) != expected_props:
                raise ValueError(
                    "existing direct attribution conflicts with the incoming energy plan"
                )

        factor_node_id = stable_id("EmissionFactor", accepted.factor_source_id)
        common_props = {
            "recordId": accepted.record_id,
            "sourceIdentity": accepted.source_identity,
            "sourceRecordId": accepted.source_record_id,
            "evidenceSourceId": accepted.evidence_source_id,
            "formulaCode": accepted.formula_code,
            "recordedScope": accepted.recorded_scope,
            "requestedScope": accepted.requested_scope,
            "systemBoundary": accepted.system_boundary,
            "isValidZero": accepted.is_valid_zero,
            "measured": accepted.measured,
            "dataProvenance": accepted.data_provenance,
        }
        consumption_props = {
            **common_props,
            "energyCarrierId": accepted.energy_carrier_id,
            "quantityId": accepted.quantity_id,
            "factorId": factor_node_id,
            "factorSourceId": accepted.factor_source_id,
            "factorAliasId": accepted.factor_id,
            "emissionId": accepted.emission_id,
            "attributionMode": plan.attribution_mode,
        }
        if plan.attribution_mode == "direct":
            consumption_props["productTargetId"] = accepted.product_target_id
        elif plan.attribution_mode == "allocated":
            consumption_props["allocationSetId"] = allocation_set_id
            consumption_props["allocationBasis"] = allocation_basis
            consumption_props["unattributedFraction"] = unattributed_fraction
        quantity_props = {
            **common_props,
            "consumptionId": accepted.consumption_id,
            "orderedRawOperands": [
                json.dumps(
                    asdict(operand),
                    ensure_ascii=False,
                    sort_keys=True,
                    separators=(",", ":"),
                )
                for operand in accepted.operands
            ],
            "conversionSteps": [
                json.dumps(
                    asdict(step),
                    ensure_ascii=False,
                    sort_keys=True,
                    separators=(",", ":"),
                )
                for step in accepted.conversion_steps
            ],
            "quantityValue": accepted.quantity_value,
            "quantityUnit": accepted.quantity_unit,
            "quantityConversionFactor": accepted.quantity_conversion_factor,
        }
        factor_props = {
            "factorId": factor_node_id,
            "factorSourceId": accepted.factor_source_id,
            "keyword": factor.keyword,
            "factorValue": accepted.factor_value,
            "factorUnit": factor.factor_unit,
            "factorDenominator": accepted.factor_denominator,
            "originalFactorValue": factor.factor_value,
            "originalFactorUnit": factor.original_factor_unit,
            "source": factor.source,
            "sourceRowId": factor.source_row_id,
            "matchStatus": factor.match_status,
            "confidence": factor.confidence,
            "denominatorUnit": factor.denominator_unit,
            "normalizedFactorValue": factor.normalized_factor_value,
            "normalizedDenominator": factor.normalized_denominator,
            "geography": factor.geography,
            "year": factor.year,
            "systemBoundary": factor.system_boundary,
            "sourceReference": factor.source_reference,
            "proxyStatus": factor.proxy_status,
            "validationStatus": factor.validation_status,
        }
        emission_props = {
            **common_props,
            "consumptionId": accepted.consumption_id,
            "quantityId": accepted.quantity_id,
            "factorId": factor_node_id,
            "factorSourceId": accepted.factor_source_id,
            "factorAliasId": accepted.factor_id,
            "quantityValue": accepted.quantity_value,
            "quantityUnit": accepted.quantity_unit,
            "factorValue": accepted.factor_value,
            "factorDenominator": accepted.factor_denominator,
            "emissionValue": accepted.emission_value,
            "emissionUnit": accepted.emission_unit,
        }

        staged_graph = deepcopy(self.graph)
        staged_graph.add_node(
            accepted.consumption_id, ["EnergyConsumption"], consumption_props
        )
        staged_graph.add_node(
            accepted.quantity_id, ["ConsumptionQuantity"], quantity_props
        )
        staged_graph.add_node(factor_node_id, ["EmissionFactor"], factor_props)
        carrier_props = {"carrierId": accepted.energy_carrier_id}
        if str(accepted.energy_carrier_key or "").strip():
            # The id is stable_id("EnergyCarrier", key), so recording the key
            # names the carrier without asserting anything the hash does not
            # already commit to.
            carrier_props["name"] = accepted.energy_carrier_key
        staged_graph.add_node(
            accepted.energy_carrier_id,
            ["EnergyCarrier"],
            carrier_props,
        )
        staged_graph.add_node(
            accepted.emission_id, ["CarbonEmission"], emission_props
        )

        def occurrence(relation: str, target_id: str, *identity: object) -> str:
            return stable_id(
                "EnergyCalculationRelationship",
                accepted.consumption_id,
                relation,
                target_id,
                *identity,
            )

        staged_graph.add_edge(
            accepted.consumption_id,
            "hasQuantity",
            accepted.quantity_id,
            occurrence_id=occurrence("hasQuantity", accepted.quantity_id),
        )
        staged_graph.add_edge(
            accepted.consumption_id,
            "hasFactor",
            factor_node_id,
            occurrence_id=occurrence("hasFactor", factor_node_id),
        )
        staged_graph.add_edge(
            accepted.consumption_id,
            "ofCarrier",
            accepted.energy_carrier_id,
            occurrence_id=occurrence("ofCarrier", accepted.energy_carrier_id),
        )
        staged_graph.add_edge(
            accepted.emission_id,
            "hasCarbonDriver",
            accepted.consumption_id,
            occurrence_id=stable_id(
                "EnergyCalculationRelationship",
                accepted.emission_id,
                "hasCarbonDriver",
                accepted.consumption_id,
            ),
        )
        if plan.attribution_mode == "direct":
            assert accepted.product_target_id is not None
            staged_graph.add_edge(
                accepted.consumption_id,
                "recordedForObject",
                accepted.product_target_id,
                planned_product_edge_props[accepted.product_target_id],
                occurrence_id=occurrence(
                    "recordedForObject", accepted.product_target_id, "direct"
                ),
            )
        elif plan.attribution_mode == "allocated":
            for allocation in allocations:
                staged_graph.add_edge(
                    accepted.consumption_id,
                    "recordedForObject",
                    allocation.target_component_id,
                    planned_product_edge_props[allocation.target_component_id],
                    occurrence_id=occurrence(
                        "recordedForObject",
                        allocation.target_component_id,
                        allocation_set_id,
                        allocation.evidence_record_id,
                    ),
                )
        context_props = {
            "sourceRecordId": accepted.source_record_id,
            "evidenceSourceId": accepted.evidence_source_id,
        }
        for process_node_id in process_node_ids:
            staged_graph.add_edge(
                accepted.consumption_id,
                "associatedWithProcess",
                process_node_id,
                context_props,
                occurrence_id=occurrence(
                    "associatedWithProcess", process_node_id
                ),
            )
        for resource_node_id in resource_node_ids:
            staged_graph.add_edge(
                accepted.consumption_id,
                "recordedForResource",
                resource_node_id,
                context_props,
                occurrence_id=occurrence("recordedForResource", resource_node_id),
            )
        self._commit_graph(self.graph, staged_graph)

    @staticmethod
    def validate_consumption_candidates(
        candidates: Iterable[ConsumptionCandidate],
    ) -> ValidationResultSet:
        """Validate canonical candidates without reading or mutating graph state."""
        return validate_candidates(candidates)

    def _issue(self, issue_type: str, comp: ComponentRecord, message: str) -> None:
        self.issues.append(
            {
                "issueType": issue_type,
                "globalId": comp.global_id,
                "ifcType": comp.ifc_type,
                "name": comp.name,
                "message": message,
            }
        )

    def build_stats(self) -> Dict[str, Any]:
        issue_counts = Counter(row.get("issueType", "") for row in self.issues)
        label_counts = Counter(
            label
            for node in self.graph.nodes.values()
            for label in node.get("labels", ())
        )
        edge_counts = Counter(edge["type"] for edge in self.graph.edges)
        return {
            "ifcFile": str(self.ifc_path),
            "moduleName": self.module_name,
            "factorLibraryLoadError": self.factor_library.load_error,
            "summary": dict(self.stats),
            "issueCounts": dict(issue_counts.most_common()),
            "nodeCount": len(self.graph.nodes),
            "edgeCount": len(self.graph.edges),
            "nodeLabelCounts": dict(label_counts.most_common()),
            "edgeTypeCounts": dict(edge_counts.most_common()),
        }

    def export_issues_csv(self, path: Path) -> None:
        fieldnames = ["issueType", "globalId", "ifcType", "name", "message", "componentGlobalId", "row"]
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
            writer.writeheader()
            for row in self.issues:
                writer.writerow(row)


def main(argv: list[str] | None = None) -> int:
    """Delegate the retired builder CLI to the fail-closed release entry point."""

    # Import lazily: the release runner reuses the builder class defined above.
    from dm2c_m23_canonical_release import main as release_main

    return release_main(argv)


if __name__ == "__main__":
    raise SystemExit(main())
