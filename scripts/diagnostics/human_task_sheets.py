# -*- coding: utf-8 -*-
"""D5: task sheets for the two author-run studies requested by XL15 and XL16.

build  -> blind annotation sheet (40 questions) + manual Cypher baseline sheet (30
          questions), each with a README tab, plus the sealed answer keys.
score  -> Cohen's kappa per slot for a filled annotation sheet, or strict scoring
          and timing for a filled Cypher sheet.

Usage
  python scripts/diagnostics/human_task_sheets.py build
  python scripts/diagnostics/human_task_sheets.py score-annotation <filled.xlsx>
  python scripts/diagnostics/human_task_sheets.py score-cypher <filled.xlsx>
"""
from __future__ import annotations

import json
import random
import statistics
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
RE_DIR = ROOT / "outputs/research_experiments"
BENCH = RE_DIR / "e4_benchmark/frozen_20260803_180240_v16_typea1_en_realloc/e4_balanced_benchmark_150_human_reviewed_v16_typea1_en.jsonl"
KG_JSON = RE_DIR / "multi_module_20260928/releases/typea_clean/multigranular_carbon_kg.json"
OUT = RE_DIR / "diagnostics_20260930/D5_human_tasks"
SEED = 20260930
STATUSES = ["executable", "incomplete_path", "empty_result", "unresolved_target", "clarification_required"]
PERSPECTIVES = ["product", "material", "process"]
OPERATIONS = ["value", "rank", "compare", "explain", "trace"]
TOL = 1e-4

STATUS_DEFS = {
    "executable": "the graph holds every record the question needs; a number can be returned",
    "incomplete_path": "the requested chain exists only in part, so a full answer cannot be traced (for example the largest record has no supporting quantity or factor row)",
    "empty_result": "the request is well formed but no record satisfies it (for example a material that has no accepted record in this module)",
    "unresolved_target": "a named component, material, process or identifier cannot be resolved to one entity in the graph",
    "clarification_required": "the question withholds the target, source or grouping dimension, or asks the system to choose it",
}
PERSPECTIVE_DEFS = {
    "product": "carbon attributed to building components (BIM items), including allocated factory energy",
    "material": "carbon of material records grouped or traced by material family",
    "process": "carbon of factory energy or manufacturing process records, including the part not attributed to any component",
}
OPERATION_DEFS = {
    "value": "a total, subtotal or count",
    "rank": "an ordering or the largest / smallest item",
    "compare": "two or more items set against each other",
    "explain": "the calculation evidence behind a value (quantity, factor, source)",
    "trace": "the full chain of records from component to factor",
}


def read_jsonl(path: Path):
    return [json.loads(l) for l in path.read_text(encoding="utf-8").splitlines() if l.strip()]


def stratified(rows, n, seed):
    rng = random.Random(seed)
    by_status = defaultdict(list)
    for r in rows:
        by_status[r["expected"]["status"]].append(r)
    total = len(rows)
    quota = {s: round(n * len(by_status[s]) / total) for s in STATUSES}
    while sum(quota.values()) > n:
        quota[max(quota, key=quota.get)] -= 1
    while sum(quota.values()) < n:
        quota["executable"] += 1
    chosen = []
    for s in STATUSES:
        pool = sorted(by_status[s], key=lambda r: r["case_id"])
        rng.shuffle(pool)
        # spread difficulty and perspective inside the status stratum
        pool.sort(key=lambda r: (r["difficulty"], r["expected"]["perspective"]))
        step = len(pool) / quota[s]
        chosen.extend(pool[int(i * step)] for i in range(quota[s]))
    rng.shuffle(chosen)
    return chosen, quota


def component_names(kg, ids):
    return [kg.get(i, i) for i in ids]


def load_component_names():
    data = json.loads(KG_JSON.read_text(encoding="utf-8"))
    names = {}
    for node in data.get("nodes", []):
        if "BuildingComponent" in (node.get("labels") or [node.get("label")]) or str(node.get("id", "")).startswith("BuildingComponent:"):
            props = node.get("props") or node.get("properties") or {}
            names[node["id"]] = f"{props.get('name', '')} [{props.get('ifcClass', '')}]"
    return names


def style_header(ws, widths):
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDDDDD")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def add_readme(wb, title, lines):
    ws = wb.create_sheet("README", 0)
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=13)
    for i, line in enumerate(lines, 3):
        ws.cell(row=i, column=1, value=line).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 120


def build() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    rows = read_jsonl(BENCH)
    names = load_component_names()

    # ---------------------------------------------------------------- annotation sheet
    ann, quota_a = stratified(rows, 40, SEED)
    wb = Workbook()
    ws = wb.active
    ws.title = "annotation"
    ws.append(["case_id", "question", "selected_components (shown to the user in the viewer)", "your_perspective", "your_operation", "your_status", "notes"])
    for r in ann:
        ws.append([r["case_id"], r["question"], "\n".join(component_names(names, r["selected_component_ids"])) or "(none: project scope)", "", "", "", ""])
    style_header(ws, {"A": 34, "B": 70, "C": 48, "D": 16, "E": 16, "F": 22, "G": 30})
    for dv_cols, values in ((["D"], PERSPECTIVES), (["E"], OPERATIONS), (["F"], STATUSES)):
        dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        for col in dv_cols:
            dv.add(f"{col}2:{col}{len(ann) + 1}")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    add_readme(wb, "Blind annotation of 40 benchmark questions (second annotator)", [
        "Purpose. Measure inter-annotator agreement on the three reference labels of the benchmark. Work alone, without the system, the graph, or the reference answers. Do not discuss questions with the first annotator until both sheets are sealed.",
        "Task. For each question give the perspective, the operation and the boundary state that a correct answer must have, using the drop-down lists. Use notes only to record genuine ambiguity.",
        "Perspective: " + "; ".join(f"{k} = {v}" for k, v in PERSPECTIVE_DEFS.items()),
        "Operation: " + "; ".join(f"{k} = {v}" for k, v in OPERATION_DEFS.items()),
        "Boundary state: " + "; ".join(f"{k} = {v}" for k, v in STATUS_DEFS.items()),
        "Context you may use. The selected components column lists what the user had highlighted in the viewer when asking. You may open the Type A module in the viewer and the material and factory-energy tables in `inputs/case_study`, because a real user sees those; the boundary state depends on what the graph holds (for example whether a named material has any accepted record).",
        f"Sampling. 40 of 150 questions, stratified by boundary state ({', '.join(f'{k} {v}' for k, v in quota_a.items())}), seed {SEED}. Order is random.",
        "Scoring. `python scripts/diagnostics/human_task_sheets.py score-annotation <this file>` reports Cohen's kappa and raw agreement per slot against the sealed key `annotation_key_40.json`, plus a confusion table for disagreements.",
    ])
    wb.save(OUT / "annotation_sheet_40.xlsx")
    (OUT / "annotation_key_40.json").write_text(json.dumps([{"case_id": r["case_id"], "perspective": r["expected"]["perspective"], "operation": r["expected"]["operation"], "status": r["expected"]["status"]} for r in ann], indent=2), encoding="utf-8")

    # ---------------------------------------------------------------- manual Cypher sheet
    cy, quota_c = stratified(rows, 30, SEED + 1)
    wb = Workbook()
    ws = wb.active
    ws.title = "cypher_baseline"
    ws.append(["case_id", "question", "selected_components", "selected_component_ids", "start_time (hh:mm)", "end_time (hh:mm)", "minutes", "cypher_query (final)", "attempts", "reported_status", "reported_perspective", "reported_operation", "reported_numbers (label=value per line)", "notes"])
    for r in cy:
        ws.append([r["case_id"], r["question"], "\n".join(component_names(names, r["selected_component_ids"])) or "(none: project scope)", "\n".join(r["selected_component_ids"]), "", "", "", "", "", "", "", "", "", ""])
    style_header(ws, {"A": 34, "B": 60, "C": 40, "D": 40, "E": 12, "F": 12, "G": 9, "H": 70, "I": 9, "J": 22, "K": 16, "L": 16, "M": 40, "N": 30})
    for col, values in (("J", STATUSES), ("K", PERSPECTIVES), ("L", OPERATIONS)):
        dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{len(cy) + 1}")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    add_readme(wb, "Manual Cypher baseline on 30 benchmark questions (expert with graph access, no compiler)", [
        "Purpose. Establish what a domain expert who can query the graph directly achieves on the same questions, in accuracy and in time, as the human reference point for the compiler (XL15).",
        "Set-up. Neo4j 5.x, empty database. Load the Type A release script `outputs/research_experiments/multi_module_20260928/releases/typea_clean/multigranular_carbon_kg.cypher` (22,833 lines, MERGE statements, one uniqueness constraint on DM2CEntity.id) with `cypher-shell -u neo4j -p <password> -f multigranular_carbon_kg.cypher`. Loading takes a few minutes. Check: MATCH (n) RETURN count(n) must return 2993 and MATCH ()-[r]->() RETURN count(r) must return 5613.",
        "Schema. Node labels and counts: BuildingComponent 321, ComponentType 169, IfcMaterial 28, DesignQuantity 1912, MaterialConsumption 139, EnergyConsumption 16, ConsumptionQuantity 155, CarbonEmission 155, EmissionFactor 8, EnergyCarrier 2, ManufacturingActivity 61, ManufacturingResource 10, ProductionStage 12, ManufacturingProcessTemplate 3, ProductionBatch 1, ModularUnit 1. Relationship types: containsComponent, hasComponentType, hasMaterial, hasDesignQuantity, recordedForObject (consumption to component; `allocated` flag on energy allocation rows), ofMaterial, ofCarrier, recordedForResource, hasCarbonDriver, derivedFrom, hasQuantity, hasFactor, associatedWithProcess, hasActivity, usesResource, hasStage, directlyPrecedes, manufacturedBy, hasProcessTemplate, produces. Carbon values sit on CarbonEmission nodes (properties emissionValue, emissionUnit, quantityValue, factorValue, formulaCode) with derivedFrom to the MaterialConsumption or EnergyConsumption record, hasQuantity to the ConsumptionQuantity and hasFactor to the EmissionFactor. Energy allocation to components is carried by recordedForObject edges with allocated = true, allocatedFraction and allocationBasis = mass; the EnergyConsumption node holds unattributedFraction.",
        "Before the session. Spend up to 60 minutes exploring the schema freely; this time is logged separately and not counted per question. Do not read the benchmark reference answers or any system output.",
        "Per question. Record start and end clock time, write the final Cypher query, the number of query attempts before you accepted a result, the boundary state you conclude, the perspective and operation you answered under, and every number you report as label=value lines (kgCO2e values to at least four significant figures). If the question cannot be answered from the graph, record the matching state and leave numbers empty. Do not spend more than 20 minutes on one question; record a timeout in notes.",
        f"Sampling. 30 of 150 questions, stratified by boundary state ({', '.join(f'{k} {v}' for k, v in quota_c.items())}), seed {SEED + 1}. Order is random.",
        "Scoring. `python scripts/diagnostics/human_task_sheets.py score-cypher <this file>` applies the benchmark rules: state match, key-free numeric match within relative tolerance 1e-4 against the sealed key `manual_cypher_key_30.json`, slot match, strict pass, plus median and total minutes.",
    ])
    wb.save(OUT / "manual_cypher_sheet_30.xlsx")
    (OUT / "manual_cypher_key_30.json").write_text(json.dumps([{"case_id": r["case_id"], "expected": r["expected"], "difficulty": r["difficulty"], "scope": r["scope"]} for r in cy], indent=2, ensure_ascii=False), encoding="utf-8")

    summary = {"seed": SEED, "annotation_quota": quota_a, "cypher_quota": quota_c,
               "annotation_case_ids": [r["case_id"] for r in ann], "cypher_case_ids": [r["case_id"] for r in cy],
               "overlap": len({r["case_id"] for r in ann} & {r["case_id"] for r in cy})}
    (OUT / "sampling_summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print(json.dumps({k: v for k, v in summary.items() if not k.endswith("_ids")}, indent=1))


def cohen_kappa(pairs):
    n = len(pairs)
    if not n:
        return None
    po = sum(1 for a, b in pairs if a == b) / n
    ca, cb = Counter(a for a, _ in pairs), Counter(b for _, b in pairs)
    pe = sum(ca[k] * cb[k] for k in set(ca) | set(cb)) / n ** 2
    return (po - pe) / (1 - pe) if pe < 1 else 1.0


def score_annotation(path: Path) -> None:
    key = {r["case_id"]: r for r in json.loads((OUT / "annotation_key_40.json").read_text(encoding="utf-8"))}
    ws = load_workbook(path)["annotation"]
    out = {}
    for slot, col in (("perspective", 3), ("operation", 4), ("status", 5)):
        pairs, confusion = [], Counter()
        for row in ws.iter_rows(min_row=2, values_only=True):
            cid, label = row[0], (row[col] or "").strip()
            if cid in key and label:
                pairs.append((key[cid][slot], label))
                if key[cid][slot] != label:
                    confusion[f"{key[cid][slot]} -> {label}"] += 1
        out[slot] = {"n": len(pairs), "agreement": sum(a == b for a, b in pairs) / len(pairs) if pairs else None, "kappa": cohen_kappa(pairs), "disagreements": dict(confusion)}
    (path.parent / (path.stem + "_kappa.json")).write_text(json.dumps(out, indent=2), encoding="utf-8")
    print(json.dumps(out, indent=1))


def parse_numbers(text):
    values = []
    for line in str(text or "").splitlines():
        if "=" in line:
            try:
                values.append(float(line.split("=", 1)[1].strip().replace(",", "")))
            except ValueError:
                pass
    return values


def close(a, b):
    return abs(a - b) <= max(TOL, abs(b) * TOL)


def score_cypher(path: Path) -> None:
    key = {r["case_id"]: r for r in json.loads((OUT / "manual_cypher_key_30.json").read_text(encoding="utf-8"))}
    ws = load_workbook(path)["cypher_baseline"]
    rows, minutes = [], []
    for row in ws.iter_rows(min_row=2, values_only=True):
        cid = row[0]
        if cid not in key:
            continue
        exp = key[cid]["expected"]
        status, persp, op = (row[9] or "").strip(), (row[10] or "").strip(), (row[11] or "").strip()
        nums = parse_numbers(row[12])
        expected_values = [v for v in (exp.get("summary") or {}).values() if isinstance(v, (int, float)) and not isinstance(v, bool)]
        numeric = None if not expected_values else all(any(close(n, e) for n in nums) for e in expected_values)
        status_ok = status == exp["status"]
        slot_ok = persp == exp["perspective"] and op == exp["operation"]
        passed = status_ok and slot_ok and (numeric is not False)
        m = row[6]
        try:
            m = float(m) if m not in (None, "") else None
        except ValueError:
            m = None
        if m is not None:
            minutes.append(m)
        rows.append({"case_id": cid, "status_match": status_ok, "slot_match": slot_ok, "numeric_match_keyfree": numeric, "passed": passed, "minutes": m, "attempts": row[8], "difficulty": key[cid]["difficulty"]})
    n = len(rows)
    summary = {"n": n, "pass_rate": sum(r["passed"] for r in rows) / n, "status_accuracy": sum(r["status_match"] for r in rows) / n,
               "slot_accuracy": sum(r["slot_match"] for r in rows) / n,
               "numeric_accuracy_keyfree": (lambda xs: sum(xs) / len(xs) if xs else None)([r["numeric_match_keyfree"] for r in rows if r["numeric_match_keyfree"] is not None]),
               "minutes_median": statistics.median(minutes) if minutes else None, "minutes_total": sum(minutes) if minutes else None,
               "by_difficulty": {d: sum(r["passed"] for r in rows if r["difficulty"] == d) / max(1, sum(r["difficulty"] == d for r in rows)) for d in ("single_hop", "multi_hop", "batch_allocation")},
               "rows": rows}
    (path.parent / (path.stem + "_scores.json")).write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print(json.dumps({k: v for k, v in summary.items() if k != "rows"}, indent=1))


if __name__ == "__main__":
    cmd = sys.argv[1] if len(sys.argv) > 1 else "build"
    if cmd == "build":
        build()
    elif cmd == "score-annotation":
        score_annotation(Path(sys.argv[2]))
    elif cmd == "score-cypher":
        score_cypher(Path(sys.argv[2]))
    else:
        raise SystemExit(__doc__)
