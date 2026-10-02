# -*- coding: utf-8 -*-
"""Render the XL14 to XL16 material.

  python scripts/diagnostics/xl1416_docs.py appendix    -> 附录A_基准与实验设置与诊断_20260930.docx
  python scripts/diagnostics/xl1416_docs.py manuscript  -> 原文修改稿_XL10-16_20260930.docx (red changes on the XL10-13 revision)
  python scripts/diagnostics/xl1416_docs.py all
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts/multi_module"))
sys.path.insert(0, str(Path(__file__).resolve().parent))
from mm_apply_to_manuscript import Editor  # noqa: E402
from mm_docx import BLACK, fill_table, paragraph_after, style_run  # noqa: E402
from xl1416_content import build_content  # noqa: E402

APPENDIX_OUT = ROOT / "附录A_基准与实验设置与诊断_20260930.docx"
SRC = ROOT / "原文修改稿_XL10-13_20260929.docx"
MANUSCRIPT_OUT = ROOT / "原文修改稿_XL10-16_20260930.docx"
DATA_NOTE = "数据来源 outputs/research_experiments/diagnostics_20260930/D3_report/diagnostics.md；脚本 scripts/diagnostics/xl1416_docs.py 可重新生成。"


def _code_para(doc, text, size=9):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(18)
    p.paragraph_format.space_after = Pt(4)
    run = style_run(p.add_run(text), size=size, font="Consolas")
    return p


def build_appendix(C) -> Path:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    style_run(doc.add_paragraph().add_run(C.appendix["title"]), bold=True, size=12)
    doc.add_paragraph(C.appendix["intro"])
    for section in C.appendix["sections"]:
        style_run(doc.add_paragraph().add_run(section["heading"]), bold=True)
        for block in section["blocks"]:
            kind, payload = block
            if kind == "para":
                doc.add_paragraph(payload)
            elif kind == "code":
                _code_para(doc, payload)
            elif kind == "numbered":
                for i, item in enumerate(payload, 1):
                    p = doc.add_paragraph(f"{i}. {item}")
                    p.paragraph_format.left_indent = Pt(18)
                    p.paragraph_format.space_after = Pt(2)
            elif kind == "table":
                tbl = payload
                doc.add_paragraph(tbl.caption)
                fill_table(doc.add_table(rows=1, cols=len(tbl.headers)), tbl.headers, tbl.rows, color=BLACK, size=8)
                if tbl.note:
                    style_run(doc.add_paragraph().add_run(tbl.note), size=9)
                else:
                    doc.add_paragraph()
    doc.save(APPENDIX_OUT)
    return APPENDIX_OUT


def apply_manuscript(C) -> Path:
    L, M = C.labels, C.main
    doc = Document(SRC)
    ed = Editor(doc)

    # ------------------------------------------------------------ 4.3.1: measures table and pointer to Appendix A
    p431 = ed.find(startswith="The benchmark was prepared from the type-A module knowledge graph")
    cur = ed.red_paragraph_after(p431, M["sec431_para"],
                                 note=f"[意见 14、15、2] 新增指标定义段与 {L.tab_measures}（评价指标），并指向 {L.app}。原 {L.tab_measures}（措辞敏感性）顺延为 {L.tab_phrasing_new}。" + DATA_NOTE)
    cur = ed.caption_after(cur, M["tab_measures"].caption)
    ed.table_after(cur, M["tab_measures"])
    for p in (ed.find(contains="Table 4 reports the outcome of each"), ed.find(startswith="Table 4. Per-question phrasing sensitivity")):
        ed.replace_sentence(p, M["table4_old"], M["table4_new"], f"原 {M['table4_old']} 顺延为 {M['table4_new']}，因 4.3.1 新增 {L.tab_measures}。")

    # ------------------------------------------------------------ 4.3.2: figures, labels and failure classes
    p432 = ed.find(startswith="This experiment examined whether the pipeline could return correct answers")
    ed.replace_sentence(p432, M["sec432_old_sd"], M["sec432_new_sd"],
                        f"[意见 15、16] 三次运行分别为 {', '.join(f'{100 * x:.1f}%' for x in C.facts['pass_per_run'])}，均值 81.3%，总体标准差 1.4（原文 0.8 与日志不符），并加 95% bootstrap 置信区间。口径标注为 {L.tab_measures} 的 strict pass rate。")
    ed.replace_sentence(p432, M["sec432_old_boundary"], M["sec432_new_boundary"],
                        "[意见 2、15] 92.8% 为 status accuracy，补边界题的 strict pass rate 88.9%，使 0.6×76.3% + 0.4×88.9% = 81.3%。")
    ed.append_text(p432, M["sec432_append"],
                   f"[意见 16] 错误分类与程序级槽位 F1 的结论句，明细在 {L.app}.8、{L.app}.9。数据：D3_report/diagnostics.md。")

    # ------------------------------------------------------------ 4.3.4: heading, protocol, paired tests, ablation
    ed.replace_sentence(ed.find(startswith="4.3.4 Baseline comparison"), M["sec434_heading_old"], M["sec434_heading_new"], "[意见 16] 小节标题加 ablation，新增消融段落。")
    p_proto = ed.find(startswith="The purpose of this experiment is to isolate the contribution of each architectural layer")
    ed.replace_sentence(p_proto, M["sec434_old_tol"], M["sec434_new_tol"],
                        f"[意见 15] 基线 key-free 判分实际使用相对 1e-3 或绝对 0.01 的容差（tmp/build_keyfree_comparison_v16.py），原文 10⁻⁴ 与实现不符，改为实际值并指向 {L.app}.6。")
    p_full = ed.find(startswith="The full CarbonQL system achieved 82.6")
    ed.replace_sentence(p_full, M["sec434_old_full"], M["sec434_new_full"],
                        f"[意见 2、16] 同一 90 题下 {L.tab_measures} 口径的 strict pass 76.3% 与 82.6% 的关系，并加按题配对的 McNemar 精确检验。数据：D3_report/diagnostics.md，键 keyfree_numeric_90。")
    ed.replace_sentence(p_full, M["sec434_old_remaining"], M["sec434_new_remaining"],
                        f"[意见 16] 失败归因：450 条回答中执行器错误 0 条，指向 {L.app}.8。")
    p_last = ed.find(startswith="These results distinguish the contribution of each system layer")
    ed.red_paragraph_after(p_last, M["sec434_ablation_para"],
                           note=f"[意见 16] 新增消融段：编译器四阶段（V1–V4）在同 150 题上的配对结果、去 LLM 编译器的规则编译器、两项控制策略。明细与表在 {L.app}.7、{L.app}.10。" + DATA_NOTE)

    # ------------------------------------------------------------ 5.2 limitations
    p52 = ed.find(startswith="The evidence comes from one project and one factory")
    ed.append_text(p52, M["sec52_append"], f"[意见 15] 局限补一句：150 题为开发集、留出集为唯一未见措辞测试、边界状态为图属性而非人工可判。数据：{L.app}.2、{L.app}.3。")

    doc.save(MANUSCRIPT_OUT)
    return MANUSCRIPT_OUT


def main(argv=None) -> None:
    what = (argv or sys.argv[1:] or ["all"])[0]
    C = build_content()
    if what in ("appendix", "all"):
        print("written:", build_appendix(C))
    if what in ("manuscript", "all"):
        print("written:", apply_manuscript(C))


if __name__ == "__main__":
    main()
