# -*- coding: utf-8 -*-
"""Stand-alone Appendix C (verification, sensitivity and cost of the multi-module case)."""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from mm_common import ROOT  # noqa: E402
from mm_content import build_content  # noqa: E402
from mm_docx import BLACK, fill_table, style_run  # noqa: E402

OUT = ROOT / "附录C_多模块验证与开销_20260929.docx"


def main() -> None:
    C = build_content()
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)

    style_run(doc.add_paragraph().add_run(C.appendix["title"]), bold=True, size=12)
    doc.add_paragraph(C.appendix["intro"])
    for section in C.appendix["sections"]:
        style_run(doc.add_paragraph().add_run(section["heading"]), bold=True)
        for block in section["blocks"]:
            if block[0] == "para":
                doc.add_paragraph(block[1])
            elif block[0] == "table":
                tbl = block[1]
                doc.add_paragraph(tbl.caption)
                fill_table(doc.add_table(rows=1, cols=len(tbl.headers)), tbl.headers, tbl.rows, color=BLACK, size=9)
                if tbl.note:
                    style_run(doc.add_paragraph().add_run(tbl.note), size=9)
                else:
                    doc.add_paragraph()
    doc.save(OUT)
    print("written:", OUT)


if __name__ == "__main__":
    main()
