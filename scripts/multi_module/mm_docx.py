# -*- coding: utf-8 -*-
"""Shared Word rendering helpers for the multi-module manuscript documents."""
from __future__ import annotations

import re
from copy import deepcopy

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph

RED = RGBColor(0xC0, 0x00, 0x00)
GREY = RGBColor(0x80, 0x80, 0x80)
BLACK = RGBColor(0x00, 0x00, 0x00)

_TOKEN = re.compile(r"([A-Za-z\u03A3\u2211]+)(?:_\{([^}]*)\})?(?:\^\{([^}]*)\})?")


def style_run(run, color=None, bold=None, italic=None, size=None, font=None):
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if size is not None:
        run.font.size = Pt(size)
    if font is not None:
        run.font.name = font
        rpr = run._r.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        rfonts.set(qn("w:eastAsia"), font)
    return run


_MATH_TOKEN = re.compile(r"([A-Za-z\u03A3\u2211]+)(_\{[^}]*\})?(\^\{[^}]*\})?")


def add_math_runs(p, text, color=BLACK, size=None):
    """Render X_{sub}^{sup} markup as italic base plus sub/superscript runs.

    Text outside the markup is emitted as plain (non-italic) runs, one run per stretch,
    so prose paragraphs that mention a symbol keep normal typography.
    """
    runs = []
    plain = []

    def flush():
        if plain:
            runs.append(style_run(p.add_run("".join(plain)), color=color, size=size))
            plain.clear()

    pos = 0
    while pos < len(text):
        m = _MATH_TOKEN.match(text, pos)
        if m and (m.group(2) or m.group(3)):
            flush()
            runs.append(style_run(p.add_run(m.group(1)), color=color, italic=True, size=size))
            if m.group(2):
                r = style_run(p.add_run(m.group(2)[2:-1]), color=color, size=size)
                r.font.subscript = True
                runs.append(r)
            if m.group(3):
                r = style_run(p.add_run(m.group(3)[2:-1]), color=color, size=size)
                r.font.superscript = True
                runs.append(r)
            pos = m.end()
        else:
            plain.append(text[pos])
            pos += 1
    flush()
    return runs


def fill_table(t: DocxTable, headers, rows, color=BLACK, size=9):
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        style_run(cell.paragraphs[0].add_run(str(h)), color=color, bold=True, size=size)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            style_run(cells[i].paragraphs[0].add_run(str(v)), color=color, size=size)
    return t


# ------------------------------------------------------------------ positional insertion
def paragraph_after(anchor, template: Paragraph | None = None) -> Paragraph:
    """New empty paragraph right after `anchor` (Paragraph or Table); copies pPr from template."""
    new_p = OxmlElement("w:p")
    if template is not None and template._p.pPr is not None:
        new_p.append(deepcopy(template._p.pPr))
    element = anchor._p if isinstance(anchor, Paragraph) else anchor._tbl
    element.addnext(new_p)
    parent = anchor._parent
    return Paragraph(new_p, parent)


def paragraph_before(anchor: Paragraph, template: Paragraph | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    if template is not None and template._p.pPr is not None:
        new_p.append(deepcopy(template._p.pPr))
    anchor._p.addprevious(new_p)
    return Paragraph(new_p, anchor._parent)


def table_after(doc, anchor, headers, rows, color=BLACK, size=9) -> DocxTable:
    t = doc.add_table(rows=1, cols=len(headers))
    fill_table(t, headers, rows, color=color, size=size)
    element = anchor._p if isinstance(anchor, Paragraph) else anchor._tbl
    element.addnext(t._tbl)
    return t


def center(p):
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p
