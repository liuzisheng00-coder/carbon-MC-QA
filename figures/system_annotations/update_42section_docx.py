"""Insert the annotated system panels and functional structure into Section 4.2."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
import shutil

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path.cwd()
TARGET = ROOT / "4.2section.docx"
FIGURE_DIR = ROOT / "figures" / "system_annotations"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_width(cell, width_inches: float) -> None:
    width = Inches(width_inches)
    cell.width = width
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width.twips)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_fixed_layout(table, total_width_inches: float) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(int(Inches(total_width_inches).twips)))
    tbl_w.set(qn("w:type"), "dxa")
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "B7C9D6")


def style_table_text(table) -> None:
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    header_marker = header_properties.find(qn("w:tblHeader"))
    if header_marker is None:
        header_marker = OxmlElement("w:tblHeader")
        header_properties.append(header_marker)
    header_marker.set(qn("w:val"), "true")
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Times New Roman")
                    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Times New Roman")
                    run.font.size = Pt(8.2)
                    run.font.bold = row_index == 0
                    if row_index == 0:
                        run.font.color.rgb = RGBColor(255, 255, 255)
            if row_index == 0:
                set_cell_shading(cell, "1F4E78")


def add_structure_table_after(doc: Document, anchor_paragraph) -> None:
    caption = doc.add_paragraph()
    caption.style = doc.styles["Caption"]
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(6)
    caption.paragraph_format.space_after = Pt(4)
    caption.add_run("Table 4.4. Functional structure of the prototype system.")

    rows = [
        ["Interface mode", "Functional role", "Information source", "User-facing output"],
        ["Model", "IFC visualisation and component selection", "IFC geometry, object class, and GlobalId", "Selected-component query scope"],
        ["Product", "Component-level carbon account", "Accepted material and process atoms grouped by component", "C_mat, C_proc, known total, and IFC context"],
        ["Material", "Material-level carbon account", "IfcMaterial assignments and material carbon atoms", "Material total, component count, and status"],
        ["Process", "Manufacturing-stage carbon account", "Production stages, energy use, and process carbon atoms", "Stage total, energy carrier, and status"],
        ["Graph + Ask DM2C", "Calculation-path inspection and grounded access", "Graph relations plus project or selected-GlobalId scope", "Subgraph, answer, perspective, source scope, and evidence"],
    ]
    table = doc.add_table(rows=len(rows), cols=4)
    widths = [0.92, 1.50, 1.78, 1.55]
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.text = value
            set_cell_width(cell, widths[column_index])
    set_table_fixed_layout(table, sum(widths))
    style_table_text(table)

    anchor_paragraph._p.addnext(caption._p)
    caption._p.addnext(table._tbl)


def replace_picture(paragraph, image_path: Path, width_inches: float, alt_text: str) -> None:
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    run = paragraph.add_run()
    shape = run.add_picture(str(image_path), width=Inches(width_inches))
    shape._inline.docPr.set("descr", alt_text)
    shape._inline.docPr.set("title", alt_text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_with_next = True


def replace_in_runs(paragraph, replacements: dict[str, str]) -> None:
    for run in paragraph.runs:
        for old, new in replacements.items():
            if old in run.text:
                run.text = run.text.replace(old, new)


def main() -> None:
    if not TARGET.exists():
        raise FileNotFoundError(TARGET)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = ROOT / f"4.2section.backup_{timestamp}.docx"
    shutil.copy2(TARGET, backup)

    doc = Document(TARGET)
    paragraphs = doc.paragraphs
    if len(paragraphs) < 19:
        raise ValueError("Unexpected Section 4.2 structure")

    paragraphs[1].text = "4.2.1 System architecture and interface structure"
    paragraphs[1].style = doc.styles["Normal"]
    paragraphs[1].paragraph_format.space_before = Pt(8)
    paragraphs[1].paragraph_format.space_after = Pt(6)
    for run in paragraphs[1].runs:
        run.bold = True
        run.font.size = Pt(12)

    add_structure_table_after(doc, paragraphs[2])

    replace_picture(
        paragraphs[3],
        FIGURE_DIR / "annotated_panel_a_model_qa.png",
        5.75,
        "Panel a: selected BIM component, grounded question answering, and perspective-specific result.",
    )
    replace_picture(
        paragraphs[5],
        FIGURE_DIR / "annotated_panel_b_accounts.png",
        5.75,
        "Panel b: product, material, and process carbon account views.",
    )
    replace_picture(
        paragraphs[7],
        FIGURE_DIR / "annotated_panel_c_calculation_path.png",
        5.75,
        "Panel c: selected-component calculation path and grounded result.",
    )

    caption_updates = {
        4: "(a) IFC model selection and grounded multi-perspective queries.",
        6: "(b) Product-, material-, and process-perspective account views.",
        8: "(c) Selected-component calculation path and grounded result.",
        9: (
            "Figure 4.5. Prototype system for multi-perspective carbon information access. "
            "Panel (a) links IFC model selection with grounded queries and perspective-specific results; "
            "panel (b) consolidates product, material, and process account views; and panel (c) exposes "
            "the selected component's calculation path from component and consumption records to carbon "
            "emissions and quantities, alongside the corresponding grounded result. The screenshots "
            "demonstrate system operation rather than numerical performance."
        ),
    }
    for index, text in caption_updates.items():
        paragraphs[index].text = text
        paragraphs[index].alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraphs[index].paragraph_format.keep_with_next = index != 9

    replace_in_runs(paragraphs[12], {"Figure 4.5(e)": "Figure 4.5(c)"})
    replace_in_runs(paragraphs[15], {"Figure 4.5(c)": "Figure 4.5(b)"})
    replace_in_runs(paragraphs[18], {"Figure 4.5(d)": "Figure 4.5(b)"})

    temporary = ROOT / "4.2section.updated.tmp.docx"
    doc.save(temporary)
    temporary.replace(TARGET)
    print(f"backup={backup}")
    print(f"output={TARGET}")


if __name__ == "__main__":
    main()
